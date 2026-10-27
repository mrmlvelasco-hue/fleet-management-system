"""/fuel and /fuel/export.xlsx must respect the acting user's
organisational scope, exactly like every other vehicle-based list and
export in this app.

Found during the Stage 1 inventory for the Fuel + Org master prompt:
`list_fuel` builds its branch filter only from the `branch_id` query
param a caller happens to send -- it never consults the user's OWN
UserOrgScope rows the way `VehicleService.list()` does, so a
BRANCH-scoped user calling `/fuel` with no filter sees every branch's
fuel, not just their own. The same gap exists in `/fuel/export.xlsx`,
and `list_fuel`'s own KPI summary ignores the branch_id filter even
when one IS sent.
"""
import json

import pytest

from app.core.security.password import hash_password
from app.core.security.registry import sync_permissions
from app.modules.master_data.org.service import BranchService
from app.modules.master_data.reference.service import VehicleTypeService
from app.modules.master_data.vehicle.service import VehicleService
from app.modules.user_management.models import Permission, Role, User
from app.modules.user_management.org_scope_service import UserOrgScopeService


@pytest.fixture()
def env(db):
    sync_permissions()
    branch_a = BranchService().create(code="BR-FSC-A", name="Fuel Scope A")
    branch_b = BranchService().create(code="BR-FSC-B", name="Fuel Scope B")
    vt = VehicleTypeService().create(code="LV-FSC", name="Light",
                                     category="LIGHT")
    v_a = VehicleService().create(vehicle_type_id=vt.id, brand="Toyota",
                                  model="Vios", year=2023,
                                  branch_id=branch_a.id,
                                  conduction_number="FSC-A-1",
                                  plate_number="FSC-A1")
    v_b = VehicleService().create(vehicle_type_id=vt.id, brand="Toyota",
                                  model="Vios", year=2023,
                                  branch_id=branch_b.id,
                                  conduction_number="FSC-B-1",
                                  plate_number="FSC-B1")

    from datetime import datetime
    from app.extensions import db as _db
    from app.modules.transactions.fuel.models import FuelTransaction
    _db.session.add(FuelTransaction(
        vehicle_id=v_a.id, transaction_date=datetime(2026, 9, 1),
        litres=10, total_amount=500, reference_number="FSC-A-REF"))
    _db.session.add(FuelTransaction(
        vehicle_id=v_b.id, transaction_date=datetime(2026, 9, 1),
        litres=20, total_amount=1000, reference_number="FSC-B-REF"))
    _db.session.commit()

    role = Role(name="Fuel Scope Role")
    role.permissions = Permission.query.filter(
        Permission.code.in_(["fuel.view"])).all()
    user = User(username="fuelscoped", email="fs@e.com",
               password_hash=hash_password("secret123"), is_active=True)
    user.roles = [role]
    from app.extensions import db as _db2
    _db2.session.add_all([role, user])
    _db2.session.commit()

    UserOrgScopeService().assign(
        user.id, scope_type="BRANCH", branch_id=branch_a.id)

    return {"branch_a": branch_a, "branch_b": branch_b,
            "v_a": v_a, "v_b": v_b, "user": user}


def _token(client, username="fuelscoped"):
    r = client.post("/api/v1/auth/token",
                    json={"username": username, "password": "secret123"})
    return json.loads(r.get_data(as_text=True)).get("access_token")


def _get(client, url, token):
    r = client.get(url, headers={"Authorization": f"Bearer {token}"})
    return r.status_code, json.loads(r.get_data(as_text=True))


def test_a_branch_scoped_user_only_sees_their_own_branchs_fuel(
        db, client, env):
    status, body = _get(client, "/api/v1/fuel", _token(client))
    assert status == 200
    refs = {item["reference_number"] for item in body["items"]}
    assert refs == {"FSC-A-REF"}


def test_the_kpi_summary_on_the_list_is_also_scoped_not_fleet_wide(
        db, client, env):
    """Even when the row list is correctly scoped, the KPI tiles beside
    it must not quietly show the whole fleet's totals."""
    status, body = _get(client, "/api/v1/fuel", _token(client))
    assert status == 200
    assert body["summary"]["fills"] == 1
    assert float(body["summary"]["spend"]) == 500


def test_export_is_scoped_the_same_way_as_the_list(db, client, env):
    token = _token(client)
    r = client.get("/api/v1/fuel/export.xlsx",
                   headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(r.data))
    ws = wb.active
    plates = [row[1].value for row in ws.iter_rows(min_row=2)]
    assert env["v_a"].plate_number in plates
    assert env["v_b"].plate_number not in plates


def test_an_unscoped_global_user_still_sees_everything(db, client, env):
    """The backward-compatibility rule used everywhere else: a user
    with NO UserOrgScope rows at all is not yet opted into scoping and
    must keep seeing the whole fleet."""
    sync_permissions()
    role = Role(name="Fuel Global Role")
    role.permissions = Permission.query.filter(
        Permission.code.in_(["fuel.view"])).all()
    user = User(username="fuelglobal", email="fg@e.com",
               password_hash=hash_password("secret123"), is_active=True)
    user.roles = [role]
    from app.extensions import db as _db
    _db.session.add_all([role, user])
    _db.session.commit()

    r = client.post("/api/v1/auth/token",
                    json={"username": "fuelglobal", "password": "secret123"})
    token = json.loads(r.get_data(as_text=True)).get("access_token")
    status, body = _get(client, "/api/v1/fuel", token)
    assert status == 200
    refs = {item["reference_number"] for item in body["items"]}
    assert refs == {"FSC-A-REF", "FSC-B-REF"}
