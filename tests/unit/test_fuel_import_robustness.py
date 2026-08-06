"""Fuel import: column matching, batch tracking, exception-count
consistency.

A real sample file (177 rows, this app's own export headers used as an
import) exposed two genuine bugs:
  * Header matching was exact-string-only, so punctuation in a header
    ("Odometer (reported)") matched NOTHING, and every single reading
    imported as MISSING regardless of the actual data -- including
    this app's OWN export column names.
  * The Exceptions badge used a 30-day-scoped count while the
    Exceptions list itself had no date restriction, so an older
    unresolved flag could show "0" on the badge while still appearing
    on the list.
"""
from datetime import datetime, timedelta
from io import BytesIO

import openpyxl
import pytest

from app.modules.transactions.fuel.import_export import (
    map_headers, _norm, build_template)


def test_this_apps_own_export_headers_all_map():
    """The exact regression: punctuation in a header must not defeat
    matching against a plain-text alias."""
    headers = ["Date", "Vehicle", "Card", "Station", "Fuel Type", "Litres",
              "Price/L", "Total", "Odometer (reported)", "Odometer (used)",
              "Odometer Status", "Distance (km)", "km/L", "Cost/km",
              "Flags", "Note"]
    mapping = map_headers(headers)
    for field in ("transaction_date", "plate_number", "card_number",
                 "station", "fuel_type", "litres", "price_per_litre",
                 "total_amount", "odometer_reported"):
        assert field in mapping, f"{field} did not map from real headers"


@pytest.mark.parametrize("header,field", [
    ("Odometer (reported)", "odometer_reported"),
    ("Price/L", "price_per_litre"),
    ("Card", "card_number"),
    ("Fuel Type", "fuel_type"),
    ("Odometer Reported", "odometer_reported"),
])
def test_specific_punctuated_headers_match_their_field(header, field):
    assert _norm(header) in {_norm(a) for a in
                            __import__("app.modules.transactions.fuel."
                                      "import_export",
                                      fromlist=["COLUMN_ALIASES"])
                            .COLUMN_ALIASES[field]}


def test_unrecognized_columns_are_still_ignored_not_fatal():
    """Fixing the false-negative matching must not turn into false
    positives -- a genuinely unrelated column must still be skipped."""
    headers = ["Date", "Vehicle", "Litres", "Total", "Some Other Column"]
    mapping = map_headers(headers)
    assert "some_other_column" not in mapping
    assert len(mapping) == 4


@pytest.fixture()
def fuel_vehicle(db):
    from app.modules.master_data.reference.service import VehicleTypeService
    from app.modules.master_data.org.service import BranchService
    from app.modules.master_data.vehicle.service import VehicleService
    vt = VehicleTypeService().create(code="LV-FIR", name="Light",
                                     category="LIGHT")
    branch = BranchService().create(code="BR-FIR", name="Branch FIR")
    vehicle = VehicleService().create(
        vehicle_type_id=vt.id, brand="Toyota", model="Hilux", year=2022,
        branch_id=branch.id, conduction_number="FIR-1",
        plate_number="FIR-111")
    db.session.commit()
    return vehicle


def _workbook_with_punctuated_headers(vehicle, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fuel"
    ws.append(["Date", "Vehicle", "Fuel Type", "Litres", "Price/L", "Total",
              "Odometer (reported)"])
    ws.append(["2026-08-01 09:00", vehicle.plate_number, "Diesel", 45.5,
              62.15, 2827.83, 65430])
    wb.save(path)
    return path


def test_real_world_headers_actually_capture_the_odometer(
        db, fuel_vehicle, tmp_path):
    """End-to-end: a file using this app's own export-style headers must
    actually populate odometer_reported, not silently import as
    MISSING."""
    from app.modules.transactions.fuel.import_export import import_fuel
    from app.modules.transactions.fuel.models import FuelTransaction

    path = _workbook_with_punctuated_headers(fuel_vehicle, tmp_path / "f.xlsx")
    with open(path, "rb") as f:
        stats = import_fuel(f, dry_run=False)
    assert stats["created"] == 1

    txn = FuelTransaction.query.filter_by(vehicle_id=fuel_vehicle.id).first()
    assert txn.odometer_reported == 65430
    assert txn.odometer_status != "MISSING"


def test_imported_row_is_stamped_with_its_source_file(
        db, fuel_vehicle, tmp_path):
    """Traceability: every imported row must record which file it came
    from, who imported it, and when -- so a wrong upload can be found
    and undone."""
    from app.modules.transactions.fuel.import_export import import_fuel
    from app.modules.transactions.fuel.models import FuelTransaction

    path = _workbook_with_punctuated_headers(fuel_vehicle, tmp_path / "f.xlsx")
    with open(path, "rb") as f:
        stats = import_fuel(f, dry_run=False, filename="statement.xlsx",
                            user_id=1)
    txn = FuelTransaction.query.filter_by(vehicle_id=fuel_vehicle.id).first()
    assert txn.import_filename == "statement.xlsx"
    assert txn.imported_by == 1
    assert txn.import_batch_id == stats["import_batch_id"]
    assert txn.imported_at is not None


def test_two_uploads_get_different_batch_ids(db, fuel_vehicle, tmp_path):
    """Each upload must be independently undoable -- two imports of the
    same file must not share a batch id."""
    from app.modules.transactions.fuel.import_export import import_fuel
    path = _workbook_with_punctuated_headers(fuel_vehicle, tmp_path / "f.xlsx")
    with open(path, "rb") as f:
        s1 = import_fuel(f, dry_run=False, filename="a.xlsx")
    with open(path, "rb") as f:
        s2 = import_fuel(f, dry_run=False, filename="b.xlsx")
    assert s1["import_batch_id"] != s2["import_batch_id"]


def test_deleting_a_batch_removes_only_that_batchs_rows(
        db, fuel_vehicle, tmp_path):
    """The undo mechanism: deleting one batch must not touch a
    transaction from a different upload or a manual entry."""
    from app.modules.transactions.fuel.import_export import (
        import_fuel, delete_import_batch)
    from app.modules.transactions.fuel.models import FuelTransaction

    path = _workbook_with_punctuated_headers(fuel_vehicle, tmp_path / "f.xlsx")
    with open(path, "rb") as f:
        stats = import_fuel(f, dry_run=False, filename="bad_upload.xlsx")
    manual = FuelTransaction(vehicle_id=fuel_vehicle.id,
                            transaction_date=datetime.now(),
                            litres=10, total_amount=600, source="MANUAL")
    db.session.add(manual)
    db.session.commit()

    deleted = delete_import_batch(stats["import_batch_id"])
    assert deleted == 1
    assert FuelTransaction.query.filter_by(
        import_batch_id=stats["import_batch_id"]).count() == 0
    assert db.session.get(FuelTransaction, manual.id) is not None


def test_exceptions_badge_matches_what_the_list_actually_shows(db):
    """The other real bug: the badge previously used a 30-day-scoped
    count while the Exceptions list had no date restriction at all."""
    from app.modules.master_data.reference.service import VehicleTypeService
    from app.modules.master_data.org.service import BranchService
    from app.modules.master_data.vehicle.service import VehicleService
    from app.modules.transactions.fuel.models import FuelTransaction
    from app.modules.transactions.fuel.analytics import FuelAnalyticsService

    vt = VehicleTypeService().create(code="LV-EXC", name="Light",
                                     category="LIGHT")
    branch = BranchService().create(code="BR-EXC", name="Branch EXC")
    vehicle = VehicleService().create(
        vehicle_type_id=vt.id, brand="Ford", model="Ranger", year=2020,
        branch_id=branch.id, conduction_number="EXC-1")
    db.session.commit()

    old_txn = FuelTransaction(
        vehicle_id=vehicle.id,
        transaction_date=datetime.now() - timedelta(days=90),
        litres=40, total_amount=2000, odometer_reported=999999,
        odometer_status="SUSPECT")
    db.session.add(old_txn)
    db.session.commit()

    badge_count = FuelAnalyticsService().exceptions_count()["count"]
    list_count = FuelTransaction.query.filter(
        db.or_(FuelTransaction.anomaly_flags.isnot(None),
              FuelTransaction.odometer_status.in_(("SUSPECT", "MISSING")))
    ).count()
    assert badge_count == list_count == 1


def test_needs_review_does_not_double_count_a_row_flagged_both_ways(db):
    """A row can carry both an anomaly flag and a suspect odometer;
    summing flagged + untrusted_odometer for a single combined figure
    would double-count it."""
    from app.modules.master_data.reference.service import VehicleTypeService
    from app.modules.master_data.org.service import BranchService
    from app.modules.master_data.vehicle.service import VehicleService
    from app.modules.transactions.fuel.models import FuelTransaction
    from app.modules.transactions.fuel.analytics import FuelAnalyticsService

    vt = VehicleTypeService().create(code="LV-NR", name="Light",
                                     category="LIGHT")
    branch = BranchService().create(code="BR-NR", name="Branch NR")
    vehicle = VehicleService().create(
        vehicle_type_id=vt.id, brand="Ford", model="Ranger", year=2020,
        branch_id=branch.id, conduction_number="NR-1")
    db.session.commit()

    both = FuelTransaction(
        vehicle_id=vehicle.id, transaction_date=datetime.now(),
        litres=40, total_amount=2000, odometer_reported=999999,
        odometer_status="SUSPECT", anomaly_flags="OVER_TANK")
    db.session.add(both)
    db.session.commit()

    summary = FuelAnalyticsService().summary()
    assert summary["flagged"] == 1
    assert summary["untrusted_odometer"] == 1
    assert summary["needs_review"] == 1   # not 2


def test_template_warns_about_realistic_distance_between_fills():
    """Guidance added directly from the sample-file finding: a reading
    implying an unrealistic distance since the last fill needs a
    second look, not a silent pass-through."""
    wb = openpyxl.load_workbook(BytesIO(build_template()))
    ref = wb["Reference"]
    text = "\n".join(str(c.value) for row in ref.iter_rows() for c in row
                    if c.value)
    assert "3,000 km" in text or "3000 km" in text


def test_summary_falls_back_to_all_time_when_recent_window_is_empty(db):
    """Reported: a real 9-row import all dated 34 days ago showed
    'Fills (last 30 days): 0, Spend: P0' next to a table plainly
    showing all 9 rows -- reproduced exactly with the reported file.
    The window must fall back to all-time data rather than show a
    misleading zero next to a list that has rows in it."""
    from datetime import datetime, timedelta
    from app.modules.master_data.reference.service import VehicleTypeService
    from app.modules.master_data.org.service import BranchService
    from app.modules.master_data.vehicle.service import VehicleService
    from app.modules.transactions.fuel.models import FuelTransaction
    from app.modules.transactions.fuel.analytics import FuelAnalyticsService

    vt = VehicleTypeService().create(code="LV-ALLT", name="Light",
                                     category="LIGHT")
    branch = BranchService().create(code="BR-ALLT", name="Branch ALLT")
    vehicle = VehicleService().create(
        vehicle_type_id=vt.id, brand="Toyota", model="Vios", year=2020,
        branch_id=branch.id, conduction_number="ALLT-1")
    db.session.commit()

    old = FuelTransaction(
        vehicle_id=vehicle.id,
        transaction_date=datetime.now() - timedelta(days=34),
        litres=45, total_amount=2700, odometer_reported=10000)
    db.session.add(old)
    db.session.commit()

    summary = FuelAnalyticsService().summary()
    assert summary["is_all_time"] is True
    assert summary["fills"] == 1
    assert summary["spend"] == 2700


def test_summary_stays_scoped_to_30_days_when_recent_data_exists(db):
    """The fallback must only engage when the window is genuinely
    empty -- if there IS recent data, older transactions outside the
    window must still be excluded as before."""
    from datetime import datetime, timedelta
    from app.modules.master_data.reference.service import VehicleTypeService
    from app.modules.master_data.org.service import BranchService
    from app.modules.master_data.vehicle.service import VehicleService
    from app.modules.transactions.fuel.models import FuelTransaction
    from app.modules.transactions.fuel.analytics import FuelAnalyticsService

    vt = VehicleTypeService().create(code="LV-RECENT", name="Light",
                                     category="LIGHT")
    branch = BranchService().create(code="BR-RECENT", name="Branch RECENT")
    vehicle = VehicleService().create(
        vehicle_type_id=vt.id, brand="Toyota", model="Vios", year=2020,
        branch_id=branch.id, conduction_number="RECENT-1")
    db.session.commit()

    db.session.add(FuelTransaction(
        vehicle_id=vehicle.id,
        transaction_date=datetime.now() - timedelta(days=60),
        litres=45, total_amount=2700))
    db.session.add(FuelTransaction(
        vehicle_id=vehicle.id, transaction_date=datetime.now() - timedelta(days=5),
        litres=40, total_amount=2400))
    db.session.commit()

    summary = FuelAnalyticsService().summary()
    assert summary["is_all_time"] is False
    assert summary["fills"] == 1          # only the recent one
    assert summary["spend"] == 2400


def test_summary_empty_fleet_is_not_marked_all_time(db):
    """A fleet with zero fuel history anywhere must not claim the empty
    result is an 'all time' figure -- there is nothing to fall back
    to."""
    from app.modules.transactions.fuel.analytics import FuelAnalyticsService
    summary = FuelAnalyticsService().summary()
    assert summary["is_all_time"] is False
    assert summary["fills"] == 0
