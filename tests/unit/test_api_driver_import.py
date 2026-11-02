"""Driver / Assignee bulk import from Excel.

Same shape as the vendor/vehicle imports (tests/unit/test_api_vendor_import.py,
tests/unit/test_api_vehicle_import.py), per the client's explicit request to
"proceed in the same process of vendor and vehicle upload" for master-data
creation week:

  GET  /api/v1/drivers/import/template  -> fill-in .xlsx, with this
       install's own Branch/Department codes on a "Reference" sheet
  POST /api/v1/drivers/import           -> multipart `file`; a PREVIEW
       unless commit=1, so uploading the wrong file saves nothing

Agreed: an employee_number or license_number that already exists is
SKIPPED and reported -- an import never changes an existing record. A
row that fails validation is skipped; the rest still import. Gated on
driver.create, as creating one driver/assignee is.

No photo is ever attached by the importer -- REQUIRE_ASSIGNEE_PHOTO (and
REQUIRE_DRIVER_LICENSE, if needed) must be switched off in System
Administration for the duration of the bulk import, exactly like the
existing JSON path on POST /api/v1/drivers already assumes.
"""
import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.core.security.password import hash_password
from app.core.security.registry import sync_permissions
from app.modules.master_data.driver.import_service import TEMPLATE_COLUMNS
from app.modules.master_data.driver.models import Driver
from app.modules.master_data.driver.service import DriverService
from app.modules.master_data.org.service import (
    BranchService, DepartmentService, RegionService, AreaDistrictService,
    OrgCompanyService)
from app.modules.system_admin.models import SystemParameter
from app.modules.user_management.models import Permission, Role, User

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _param(db, code, value, group="DRIVER"):
    row = SystemParameter.query.filter_by(code=code).first()
    if row is None:
        row = SystemParameter(code=code, value=str(value),
                              data_type="BOOLEAN", group_name=group,
                              description=f"Migration switch: {code}",
                              is_active=True)
        db.session.add(row)
    else:
        row.value = str(value)
    db.session.commit()


@pytest.fixture()
def env(db):
    sync_permissions()
    # The importer never sends a photo (see import_service.py's module
    # docstring) -- bulk import depends on these being switched off for
    # the duration of the batch, exactly as a real migration would.
    _param(db, "REQUIRE_ASSIGNEE_PHOTO", "NO")
    _param(db, "REQUIRE_DRIVER_LICENSE", "NO")
    br_mun = BranchService().create(code="BR-DIMUN", name="Muntinlupa")
    BranchService().create(code="BR-DILP", name="Las Piñas")
    DepartmentService().create(code="LOG", name="Logistics",
                               branch_id=br_mun.id)
    DriverService().create(
        employee_number="EMP-OLD", first_name="Already", last_name="Here",
        branch_id=br_mun.id, assignee_type="EMPLOYEE")

    def mk(username, codes):
        role = Role(name=f"role-{username}")
        role.permissions = Permission.query.filter(
            Permission.code.in_(codes)).all()
        u = User(username=username, email=f"{username}@e.com",
                password_hash=hash_password("secret123"), is_active=True)
        u.roles = [role]
        db.session.add_all([role, u])

    mk("dimporter", ["driver.view", "driver.create"])
    mk("dviewer", ["driver.view"])
    db.session.commit()
    return {"br_mun": br_mun}


def _hdr(client, username="dimporter"):
    r = client.post("/api/v1/auth/token",
                    json={"username": username, "password": "secret123"})
    tok = json.loads(r.get_data(as_text=True))["access_token"]
    return {"Authorization": f"Bearer {tok}"}


def _sheet(rows, headers=None, title="Drivers"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    headers = headers or [c for c, _req, _note in TEMPLATE_COLUMNS]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


GOOD = {"employee_number": "EMP-101", "first_name": "Juan",
       "last_name": "Dela Cruz", "assignee_type": "driver",
       "branch_code": "br-dimun", "department_code": "log",
       "license_number": "N01-23-456789", "license_expiry": "2027-05-15",
       "license_type": "professional", "phone": "0917 000 0000",
       "email": "juan@example.ph"}


def _post(client, buf, commit=None, user="dimporter", name="drivers.xlsx"):
    data = {"file": (buf, name)}
    if commit is not None:
        data["commit"] = commit
    return client.post("/api/v1/drivers/import", headers=_hdr(client, user),
                       data=data, content_type="multipart/form-data")


def _problems(body):
    return {e["identifier"]: " | ".join(e["problems"]) for e in body["errors"]}


# ── template ────────────────────────────────────────────────────────────

def test_template_is_a_workbook_with_this_installs_codes(client, db, env):
    r = client.get("/api/v1/drivers/import/template", headers=_hdr(client))
    assert r.status_code == 200
    assert XLSX in r.headers["Content-Type"]
    assert "Driver_Assignee_Import_Template.xlsx" in r.headers[
        "Content-Disposition"]
    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames[:2] == ["Drivers", "Reference"]
    headers = [c.value for c in wb["Drivers"][1]]
    assert headers == [c for c, _r, _n in TEMPLATE_COLUMNS]
    refs = " ".join(str(c.value) for row in wb["Reference"].iter_rows()
                    for c in row if c.value)
    assert "BR-DIMUN" in refs and "BR-DILP" in refs
    assert "LOG" in refs
    assert "DRIVER" in refs and "CONSULTANT" in refs


def test_template_requires_driver_create(client, db, env):
    r = client.get("/api/v1/drivers/import/template",
                   headers=_hdr(client, "dviewer"))
    assert r.status_code == 403


def test_the_templates_own_example_row_is_never_imported(client, db, env):
    tmpl = client.get("/api/v1/drivers/import/template", headers=_hdr(client)).data
    r = _post(client, BytesIO(tmpl), commit="1")
    body = r.get_json()
    assert body["total_rows"] == 0 and body["created"] == 0
    assert Driver.query.filter(
        Driver.employee_number.like("EXAMPLE%")).count() == 0


# ── preview vs commit ───────────────────────────────────────────────────

def test_preview_is_the_default_and_saves_nothing(client, db, env):
    r = _post(client, _sheet([GOOD]))
    assert r.status_code == 200, r.get_data(as_text=True)
    body = r.get_json()
    assert body["dry_run"] is True
    assert body["created"] == 1 and body["skipped"] == 0
    assert Driver.query.filter_by(employee_number="EMP-101").count() == 0


def test_commit_zero_is_still_a_preview(client, db, env):
    _post(client, _sheet([GOOD]), commit="0")
    assert Driver.query.filter_by(employee_number="EMP-101").count() == 0


def test_commit_saves_every_field_without_a_photo(client, db, env):
    r = _post(client, _sheet([GOOD]), commit="1")
    body = r.get_json()
    assert body["dry_run"] is False and body["created"] == 1
    d = Driver.query.filter_by(employee_number="EMP-101").one()
    assert (d.first_name, d.last_name) == ("Juan", "Dela Cruz")
    assert d.assignee_type == "DRIVER"          # normalised to upper case
    assert d.license_number == "N01-23-456789"
    assert d.license_type == "professional"
    assert d.photo_attachment_id is None


def test_blank_type_defaults_to_driver_and_blank_status_to_active(
        client, db, env):
    _post(client, _sheet([{"employee_number": "EMP-102",
                           "first_name": "Plain", "last_name": "Person",
                           "branch_code": "BR-DIMUN"}]), commit="1")
    d = Driver.query.filter_by(employee_number="EMP-102").one()
    assert d.assignee_type == "DRIVER"
    assert d.status == "ACTIVE"


# ── row validation: skipped and reported, the rest still import ─────────

def test_an_existing_employee_number_is_skipped_and_untouched(client, db, env):
    r = _post(client, _sheet([{"employee_number": "emp-old",
                               "first_name": "Overwrite", "last_name": "Attempt",
                               "branch_code": "BR-DIMUN"}, GOOD]), commit="1")
    body = r.get_json()
    assert body["created"] == 1 and body["skipped"] == 1
    assert "already exists" in _problems(body)["emp-old"]
    old = Driver.query.filter_by(employee_number="EMP-OLD").one()
    assert (old.first_name, old.last_name) == ("Already", "Here")


def test_an_employee_number_repeated_in_the_same_file_imports_once(
        client, db, env):
    r = _post(client, _sheet(
        [GOOD, {**GOOD, "first_name": "Second"}]), commit="1")
    body = r.get_json()
    assert body["created"] == 1 and body["skipped"] == 1
    assert "appears more than once" in _problems(body)["EMP-101"]
    assert Driver.query.filter_by(
        employee_number="EMP-101").one().first_name == "Juan"


def test_each_bad_row_names_its_own_problems(client, db, env):
    rows = [
        {"employee_number": "EMP-201", "branch_code": "BR-DIMUN"},  # no names
        {"first_name": "No", "last_name": "Code", "branch_code": "BR-DIMUN"},
        {"employee_number": "EMP-202", "first_name": "X", "last_name": "Y",
         "branch_code": "BR-DIMUN", "assignee_type": "VENDOR"},
        {"employee_number": "EMP-203", "first_name": "X", "last_name": "Y",
         "branch_code": "NOWHERE"},
        {"employee_number": "EMP-204", "first_name": "X", "last_name": "Y",
         "branch_code": "BR-DIMUN", "email": "not-an-email"},
        {"employee_number": "EMP-205", "first_name": "X", "last_name": "Y",
         "branch_code": "BR-DIMUN", "license_expiry": "not-a-date"},
        GOOD,
    ]
    body = _post(client, _sheet(rows), commit="1").get_json()
    assert body["total_rows"] == 7
    assert body["created"] == 1 and body["skipped"] == 6
    p = _problems(body)
    assert "first_name is required" in p["EMP-201"]
    assert "employee_number is required" in p["No Code"]
    assert "assignee_type 'VENDOR'" in p["EMP-202"]
    assert "unknown branch code 'NOWHERE'" in p["EMP-203"]
    assert "email" in p["EMP-204"]
    assert "not a recognizable date" in p["EMP-205"]
    assert Driver.query.filter(
        Driver.employee_number.like("EMP-2%")).count() == 0
    assert Driver.query.filter_by(employee_number="EMP-101").count() == 1


def test_preview_reports_the_same_problems_a_commit_would(client, db, env):
    rows = [GOOD, {"employee_number": "EMP-202", "first_name": "X",
                   "last_name": "Y", "branch_code": "BR-DIMUN",
                   "assignee_type": "VENDOR"}]
    preview = _post(client, _sheet(rows)).get_json()
    assert (preview["created"], preview["skipped"]) == (1, 1)
    assert "assignee_type 'VENDOR'" in _problems(preview)["EMP-202"]


def test_blank_rows_are_ignored(client, db, env):
    body = _post(client, _sheet([{}, GOOD, {}])).get_json()
    assert body["total_rows"] == 1


def test_unknown_department_code_for_the_branch_is_reported(client, db, env):
    r = _post(client, _sheet([{**GOOD, "employee_number": "EMP-206",
                               "department_code": "NOPE"}]), commit="1")
    body = r.get_json()
    assert body["skipped"] == 1
    assert "unknown department code" in _problems(body)["EMP-206"]


# ── file-level problems ─────────────────────────────────────────────────

def test_missing_required_columns_is_a_400_naming_them(client, db, env):
    r = _post(client, _sheet([{"employee_number": "A"}],
                             headers=["employee_number", "phone"]))
    assert r.status_code == 400
    assert "first_name" in r.get_json()["message"]


def test_a_file_that_is_not_a_workbook_is_a_400(client, db, env):
    r = _post(client, BytesIO(b"not a spreadsheet"), name="drivers.xlsx")
    assert r.status_code == 400
    assert "Could not read the file" in r.get_json()["message"]


def test_no_file_is_a_400(client, db, env):
    r = client.post("/api/v1/drivers/import", headers=_hdr(client),
                    data={}, content_type="multipart/form-data")
    assert r.status_code == 400


def test_import_requires_driver_create(client, db, env):
    r = _post(client, _sheet([GOOD]), commit="1", user="dviewer")
    assert r.status_code == 403
    assert Driver.query.filter_by(employee_number="EMP-101").count() == 0


def test_a_first_sheet_with_another_name_is_still_read(client, db, env):
    body = _post(client, _sheet([GOOD], title="Sheet1")).get_json()
    assert body["created"] == 1


# ── Batch 9: Company-scoped import (API layer) ─────────────────────────

@pytest.fixture()
def scoped_env(db, env):
    region = RegionService().create(code="RG-DIAPI", name="Region DI API")
    area = AreaDistrictService().create(code="AD-DIAPI", name="Area DI API",
                                        region_id=region.id)
    company = OrgCompanyService().create(short_name="Co DI-API",
                                         region_id=region.id,
                                         area_district_id=area.id)
    scoped_branch = BranchService().create(code="BR-DIAPI-A", name="Scoped",
                                           company_id=company.id)
    db.session.commit()
    return {"company": company, "scoped_branch": scoped_branch}


def test_template_download_accepts_a_company_id_and_scopes_branch_codes(
        client, db, env, scoped_env):
    r = client.get(
        f"/api/v1/drivers/import/template"
        f"?company_id={scoped_env['company'].id}",
        headers=_hdr(client))
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.data))
    refs = " ".join(str(c.value) for row in wb["Reference"].iter_rows()
                    for c in row if c.value)
    # Scope the check to the "Valid Branch codes:" section itself --
    # the separate "Valid Department codes" section intentionally
    # lists every department together with its own branch's code, for
    # context, and is not part of this Company-scoping change.
    branch_section = refs.split("Valid Branch codes:")[1].split(
        "Valid Department codes")[0]
    assert "BR-DIAPI-A" in branch_section
    assert "BR-DIMUN" not in branch_section.split()
    assert "BR-DILP" not in branch_section.split()


def test_import_with_a_company_id_accepts_a_branch_in_that_company(
        client, db, env, scoped_env):
    row = {"employee_number": "EMP-API1", "first_name": "Ana",
          "last_name": "Reyes", "assignee_type": "employee",
          "branch_code": "BR-DIAPI-A"}
    data = {"file": (_sheet([row]), "drivers.xlsx"), "commit": "1",
            "company_id": str(scoped_env["company"].id)}
    r = client.post("/api/v1/drivers/import", headers=_hdr(client),
                    data=data, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_data(as_text=True)
    body = r.get_json()
    assert body["created"] == 1 and body["skipped"] == 0
    assert Driver.query.filter_by(employee_number="EMP-API1").count() == 1


def test_import_with_a_company_id_rejects_a_branch_outside_that_company(
        client, db, env, scoped_env):
    row = {"employee_number": "EMP-API2", "first_name": "Ana",
          "last_name": "Reyes", "assignee_type": "employee",
          "branch_code": "BR-DIMUN"}  # belongs to no company (from `env`)
    data = {"file": (_sheet([row]), "drivers.xlsx"), "commit": "1",
            "company_id": str(scoped_env["company"].id)}
    r = client.post("/api/v1/drivers/import", headers=_hdr(client),
                    data=data, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_data(as_text=True)
    body = r.get_json()
    assert body["created"] == 0 and body["skipped"] == 1
    assert "does not belong to the selected Company" in \
        body["errors"][0]["problems"][0]
    assert Driver.query.filter_by(employee_number="EMP-API2").count() == 0


def test_import_with_a_non_numeric_company_id_is_a_400(client, db, env):
    data = {"file": (_sheet([GOOD]), "drivers.xlsx"), "commit": "1",
            "company_id": "not-a-number"}
    r = client.post("/api/v1/drivers/import", headers=_hdr(client),
                    data=data, content_type="multipart/form-data")
    assert r.status_code == 400
