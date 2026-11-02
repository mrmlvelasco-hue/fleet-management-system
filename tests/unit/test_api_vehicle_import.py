"""Vehicle import over the REST API (for the React front end).

The importer itself already existed behind the Flask page at
/master/vehicles/import. React had no screen and the API had no
endpoint, so these two endpoints expose the SAME service functions
(build_template, import_vehicles) -- no second implementation of the
validation:

  GET  /api/v1/vehicles/import/template   -> the fill-in .xlsx
  POST /api/v1/vehicles/import            -> multipart `file`; nothing is
       written unless commit=1 (a preview is the default, so a wrong
       file cannot create hundreds of vehicles).

Both require vehicle.create, as the Flask page does.
"""
import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.core.security.password import hash_password
from app.core.security.registry import sync_permissions
from app.modules.master_data.org.service import (
    BranchService, RegionService, AreaDistrictService, OrgCompanyService)
from app.modules.master_data.reference.service import VehicleTypeService
from app.modules.master_data.vehicle.import_service import TEMPLATE_COLUMNS
from app.modules.master_data.vehicle.models import Vehicle
from app.modules.user_management.models import Permission, Role, User

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def env(db):
    sync_permissions()
    BranchService().create(code="BR-VI", name="Import Branch")
    VehicleTypeService().create(code="LV-VI", name="Light", category="LIGHT")

    def mk(username, codes):
        role = Role(name=f"role-{username}")
        role.permissions = Permission.query.filter(
            Permission.code.in_(codes)).all()
        u = User(username=username, email=f"{username}@e.com",
                 password_hash=hash_password("secret123"), is_active=True)
        u.roles = [role]
        db.session.add_all([role, u])

    mk("importer", ["vehicle.view", "vehicle.create"])
    mk("viewer", ["vehicle.view"])
    db.session.commit()


def _hdr(client, username="importer"):
    r = client.post("/api/v1/auth/token",
                    json={"username": username, "password": "secret123"})
    tok = json.loads(r.get_data(as_text=True))["access_token"]
    return {"Authorization": f"Bearer {tok}"}


def _sheet(rows):
    """A filled-in template: real headers, caller's rows (dicts by column)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Vehicles"
    headers = [c for c, _req, _note in TEMPLATE_COLUMNS]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


GOOD = {"conduction_number": "VI-001", "vehicle_type_code": "LV-VI",
        "branch_code": "BR-VI", "brand": "Toyota", "model": "Vios",
        "year": 2022}
BAD = {"conduction_number": "VI-002", "vehicle_type_code": "NOPE",
       "branch_code": "BR-VI", "brand": "Toyota", "model": "Vios",
       "year": 2022}


def _post(client, buf, commit=None, user="importer", name="vehicles.xlsx"):
    data = {"file": (buf, name)}
    if commit is not None:
        data["commit"] = commit
    return client.post("/api/v1/vehicles/import", headers=_hdr(client, user),
                       data=data, content_type="multipart/form-data")


# ── template ────────────────────────────────────────────────────────────

def test_template_download_is_a_workbook_with_a_reference_sheet(client, db, env):
    r = client.get("/api/v1/vehicles/import/template", headers=_hdr(client))
    assert r.status_code == 200
    assert XLSX in r.headers["Content-Type"]
    assert "Vehicle_Import_Template.xlsx" in r.headers["Content-Disposition"]
    wb = load_workbook(BytesIO(r.data))
    assert "Reference" in wb.sheetnames
    refs = " ".join(str(c.value) for ws in [wb["Reference"]]
                    for row in ws.iter_rows() for c in row if c.value)
    assert "LV-VI" in refs and "BR-VI" in refs   # this install's own codes


def test_template_requires_vehicle_create(client, db, env):
    r = client.get("/api/v1/vehicles/import/template",
                   headers=_hdr(client, "viewer"))
    assert r.status_code == 403


# ── import ──────────────────────────────────────────────────────────────

def test_preview_is_the_default_and_saves_nothing(client, db, env):
    r = _post(client, _sheet([GOOD]))
    assert r.status_code == 200, r.get_data(as_text=True)
    body = r.get_json()
    assert body["dry_run"] is True
    assert body["created"] == 1 and body["skipped"] == 0
    assert Vehicle.query.filter_by(conduction_number="VI-001").count() == 0


def test_commit_saves_the_vehicles(client, db, env):
    r = _post(client, _sheet([GOOD]), commit="1")
    assert r.status_code == 200, r.get_data(as_text=True)
    body = r.get_json()
    assert body["dry_run"] is False and body["created"] == 1
    assert Vehicle.query.filter_by(conduction_number="VI-001").count() == 1


def test_commit_zero_is_still_a_preview(client, db, env):
    r = _post(client, _sheet([GOOD]), commit="0")
    assert r.get_json()["dry_run"] is True
    assert Vehicle.query.filter_by(conduction_number="VI-001").count() == 0


# ── Batch 9: Company-scoped import (API layer) ─────────────────────────

@pytest.fixture()
def scoped_env(db, env):
    region = RegionService().create(code="RG-VIAPI", name="Region VI API")
    area = AreaDistrictService().create(code="AD-VIAPI", name="Area VI API",
                                        region_id=region.id)
    company = OrgCompanyService().create(short_name="Co VI-API",
                                         region_id=region.id,
                                         area_district_id=area.id)
    scoped_branch = BranchService().create(code="BR-VIAPI-A", name="Scoped",
                                           company_id=company.id)
    db.session.commit()
    return {"company": company, "scoped_branch": scoped_branch}


def test_template_download_accepts_a_company_id_and_scopes_branch_codes(
        client, db, env, scoped_env):
    r = client.get(
        f"/api/v1/vehicles/import/template"
        f"?company_id={scoped_env['company'].id}",
        headers=_hdr(client))
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.data))
    refs = " ".join(str(c.value) for row in wb["Reference"].iter_rows()
                    for c in row if c.value)
    assert "BR-VIAPI-A" in refs
    # "BR-VI" (unscoped, from `env`) is a substring of "BR-VIAPI-A", so
    # check for it as a whole token rather than a bare substring.
    assert "BR-VI" not in refs.split()


def test_import_with_a_company_id_accepts_a_branch_in_that_company(
        client, db, env, scoped_env):
    row = {"conduction_number": "VI-API1", "vehicle_type_code": "LV-VI",
          "branch_code": "BR-VIAPI-A", "brand": "Toyota", "model": "Vios",
          "year": 2022}
    data = {"file": (_sheet([row]), "vehicles.xlsx"), "commit": "1",
            "company_id": str(scoped_env["company"].id)}
    r = client.post("/api/v1/vehicles/import", headers=_hdr(client),
                    data=data, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_data(as_text=True)
    body = r.get_json()
    assert body["created"] == 1 and body["skipped"] == 0
    assert Vehicle.query.filter_by(conduction_number="VI-API1").count() == 1


def test_import_with_a_company_id_rejects_a_branch_outside_that_company(
        client, db, env, scoped_env):
    row = {"conduction_number": "VI-API2", "vehicle_type_code": "LV-VI",
          "branch_code": "BR-VI",  # belongs to no company (from `env`)
          "brand": "Toyota", "model": "Vios", "year": 2022}
    data = {"file": (_sheet([row]), "vehicles.xlsx"), "commit": "1",
            "company_id": str(scoped_env["company"].id)}
    r = client.post("/api/v1/vehicles/import", headers=_hdr(client),
                    data=data, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_data(as_text=True)
    body = r.get_json()
    assert body["created"] == 0 and body["skipped"] == 1
    assert "does not belong to the selected Company" in \
        body["errors"][0]["problems"][0]
    assert Vehicle.query.filter_by(conduction_number="VI-API2").count() == 0


def test_import_with_a_non_numeric_company_id_is_a_400(client, db, env):
    data = {"file": (_sheet([GOOD]), "vehicles.xlsx"), "commit": "1",
            "company_id": "not-a-number"}
    r = client.post("/api/v1/vehicles/import", headers=_hdr(client),
                    data=data, content_type="multipart/form-data")
    assert r.status_code == 400


def test_a_bad_row_is_reported_and_the_good_row_still_imports(client, db, env):
    r = _post(client, _sheet([GOOD, BAD]), commit="1")
    body = r.get_json()
    assert body["created"] == 1 and body["skipped"] == 1
    assert body["errors"][0]["row"] == 3
    assert "NOPE" in " ".join(body["errors"][0]["problems"])
    assert Vehicle.query.filter_by(conduction_number="VI-001").count() == 1
    assert Vehicle.query.filter_by(conduction_number="VI-002").count() == 0


def test_no_file_is_a_validation_error(client, db, env):
    r = client.post("/api/v1/vehicles/import", headers=_hdr(client),
                    data={}, content_type="multipart/form-data")
    assert r.status_code == 400
    assert r.get_json()["error"] == "validation"


def test_a_file_that_is_not_a_workbook_is_a_validation_error(client, db, env):
    r = _post(client, BytesIO(b"this is not an excel file"))
    assert r.status_code == 400
    assert r.get_json()["message"]


def test_a_workbook_missing_required_columns_names_them(client, db, env):
    wb = Workbook()
    wb.active.append(["brand", "model"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = _post(client, buf)
    assert r.status_code == 400
    assert "vehicle_type_code" in r.get_json()["message"]


def test_import_requires_vehicle_create(client, db, env):
    r = _post(client, _sheet([GOOD]), commit="1", user="viewer")
    assert r.status_code == 403
    assert Vehicle.query.count() == 0
