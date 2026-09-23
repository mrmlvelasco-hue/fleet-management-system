"""API endpoint exposing scripts/import_pm_task_list.py to the React
admin UI -- the same dry-run/reset/import capability the PowerShell
guide describes, reachable without SSH access to the server.

Reuses import_pm_task_list() and reset_pm_data() unchanged; this test
file does not re-verify their parsing correctness (test_import_pm_task_list.py
already does that in depth) -- only that the HTTP layer wires them up
correctly: the right permission, the right response shape, dry-run
never writing anything, and reset only ever running for real when the
caller explicitly asked for both --reset and a non-dry-run.
"""
import io
import json

import openpyxl
import pytest

from app.core.security.registry import sync_permissions
from app.core.security.password import hash_password
from app.modules.user_management.models import Permission, Role, User
from app.extensions import db


def _make_workbook_with_scope_details() -> bytes:
    """Same shape as the real VEMS export this whole importer was built
    for -- one package, one Scope Details join, matching the fixture
    already proven correct in test_import_pm_task_list.py."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Make", "Model", "Description", "Task_CD", "PM_CD",
              "WorkDescription", "Scope", "KM Reading", "KMDesc",
              "Calendar", "CALDesc", "Hourly", "HourDesc",
              "PMDescription", "Planned_CD", "Planned", "Sort"])
    ws.append([
        "Ford", "Escape", "Vehicle Preventive Maintenance Ford Escape",
        "S02-00001", "004-001-00001",
        "First 1,000 km servicing", "See Sheet Scope Details",
        "1KM", "First 1000 Km", "1MTH", "Every Month", 0, "NULL",
        "Vehicle Preventive Maintenance", "S02",
        "Vehicle Preventive Maintenance", 1])

    scope = wb.create_sheet("Scope Details")
    scope.append(["Make", "Model", "PM_CD", "Task_CD", "Line_No",
                 "Scope_Detail"])
    scope.append(["Ford", "Escape", "004-001-00001", "S02-00001", 1,
                 "Perform first 1,000 km PMS."])
    scope.append(["Ford", "Escape", "004-001-00001", "S02-00001", 2,
                 "Check tire pressure and condition."])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_workbook_missing_columns() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Make", "Model"])  # nowhere near the required columns
    ws.append(["Ford", "Escape"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def env(db):
    sync_permissions()
    db.session.commit()

    role = Role(name="PM Importer")
    role.permissions = Permission.query.filter(
        Permission.code.in_(["pmschedule.create"])).all()
    db.session.add(role)
    db.session.flush()

    admin = User(username="pmimporter", email="pmimporter@e.com",
                 password_hash=hash_password("secret123"), is_active=True)
    admin.roles.append(role)

    outsider = User(username="pmoutsider", email="pmoutsider@e.com",
                    password_hash=hash_password("secret123"),
                    is_active=True)
    db.session.add_all([admin, outsider])
    db.session.commit()
    return {"admin": admin, "outsider": outsider}


def _token(client, username):
    r = client.post("/api/v1/auth/token",
                    json={"username": username, "password": "secret123"})
    return json.loads(r.get_data(as_text=True))["access_token"]


def _upload(client, token, content, **form):
    data = {"file": (io.BytesIO(content), "pm_task_list.xlsx"), **form}
    r = client.post(
        "/api/v1/pm-task-list/import",
        data=data,
        headers={"Authorization": f"Bearer {token}"},
        content_type="multipart/form-data",
    )
    return r.status_code, json.loads(r.get_data(as_text=True))


class TestDryRun:
    def test_dry_run_returns_the_same_summary_shape_as_the_cli(self, client, env):
        content = _make_workbook_with_scope_details()
        status, body = _upload(
            client, _token(client, "pmimporter"), content, dry_run="true")

        assert status == 200
        result = body["import"]
        assert result["packages_created"] == 1
        assert result["scope_items_created"] == 2
        assert result["unrecognized_frequency_codes"] == []

    def test_dry_run_writes_nothing_to_the_database(self, client, env):
        content = _make_workbook_with_scope_details()
        _upload(client, _token(client, "pmimporter"), content, dry_run="true")

        from app.modules.maintenance_config.models import PMSchedule
        assert PMSchedule.query.count() == 0

    def test_defaults_to_dry_run_when_not_specified(self, client, env):
        # The guide's whole point: "Never skip this." A caller who
        # forgets to say which mode gets the safe one, not a silent
        # real import.
        content = _make_workbook_with_scope_details()
        status, body = _upload(client, _token(client, "pmimporter"), content)

        assert status == 200
        from app.modules.maintenance_config.models import PMSchedule
        assert PMSchedule.query.count() == 0
        assert body["import"]["packages_created"] == 1  # still reported


class TestRealImport:
    def test_writes_to_the_database_when_dry_run_is_false(self, client, env):
        content = _make_workbook_with_scope_details()
        status, body = _upload(
            client, _token(client, "pmimporter"), content, dry_run="false")

        assert status == 200
        from app.modules.maintenance_config.models import PMSchedule
        assert PMSchedule.query.count() == 1

    def test_a_malformed_file_returns_a_validation_error_not_a_500(self, client, env):
        content = _make_workbook_missing_columns()
        status, body = _upload(
            client, _token(client, "pmimporter"), content, dry_run="true")

        assert status == 400
        assert "error" in body

    def test_requires_a_file(self, client, env):
        status, body = _upload(client, _token(client, "pmimporter"), b"", dry_run="true")
        # An empty upload with no real filename content is still a
        # validation error, not a crash inside openpyxl.
        assert status == 400


class TestReset:
    def test_reset_is_never_actually_run_during_a_dry_run(self, client, env):
        from app.modules.maintenance_config.models import (
            PMSchedule, PMScopeTemplate, PMScopeItem)
        pre_existing = PMSchedule.query.count()

        content = _make_workbook_with_scope_details()
        status, body = _upload(
            client, _token(client, "pmimporter"), content,
            dry_run="true", reset="true")

        assert status == 200
        assert body["reset"] is None
        assert body["reset_skipped_dry_run"] is True
        assert PMSchedule.query.count() == pre_existing

    def test_reset_actually_clears_existing_data_before_a_real_import(
            self, client, env):
        # Seed one existing schedule the old-fashioned way, then confirm
        # --reset clears it before the new file's own row is imported.
        content = _make_workbook_with_scope_details()
        _upload(client, _token(client, "pmimporter"), content, dry_run="false")

        from app.modules.maintenance_config.models import PMSchedule
        assert PMSchedule.query.count() == 1

        status, body = _upload(
            client, _token(client, "pmimporter"), content,
            dry_run="false", reset="true")

        assert status == 200
        assert body["reset"]["pm_schedules"] == 1
        # The re-import replaces it with exactly one fresh row, not two.
        assert PMSchedule.query.count() == 1

    def test_without_reset_a_second_import_adds_alongside_the_first(
            self, client, env):
        content = _make_workbook_with_scope_details()
        _upload(client, _token(client, "pmimporter"), content, dry_run="false")
        _upload(client, _token(client, "pmimporter"), content, dry_run="false")

        from app.modules.maintenance_config.models import PMSchedule
        assert PMSchedule.query.count() == 2


class TestPermission:
    def test_requires_pmschedule_create(self, client, env):
        content = _make_workbook_with_scope_details()
        status, _ = _upload(
            client, _token(client, "pmoutsider"), content, dry_run="true")
        assert status == 403
