"""Batch 9 (2026-10-31): Driver/Assignee CSV/Excel import scoped to a
selected Company.

Same design as the vehicle importer
(tests/unit/test_vehicle_import_company_scope.py): `company_id` is an
OPTIONAL parameter to import_drivers()/build_template(). Omitted,
behaviour is unchanged (backward compatible with every existing
caller/test). Given, every branch_code in the file must resolve to a
Branch under THAT Company -- a branch that exists but belongs to a
different Company, or to none at all (an "unmapped" branch), is
rejected as a row error rather than silently accepted.
"""
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.master_data.org.service import (
    BranchService, RegionService, AreaDistrictService, OrgCompanyService)
from app.modules.master_data.driver.import_service import (
    TEMPLATE_COLUMNS, build_template, import_drivers)
from app.modules.system_admin.models import SystemParameter


def _param(db, code, value, group="DRIVER"):
    row = SystemParameter.query.filter_by(code=code).first()
    if row is None:
        row = SystemParameter(code=code, value=str(value),
                              data_type="BOOLEAN", group_name=group,
                              description=f"Migration switch: {code}",
                              is_active=True)
        db.session.add(row)
    else:
        row.value = str(value)
    db.session.commit()


def _company(suffix):
    region = RegionService().create(code=f"RG-DRV{suffix}",
                                     name=f"Region DRV {suffix}")
    area = AreaDistrictService().create(code=f"AD-DRV{suffix}",
                                        name=f"Area DRV {suffix}",
                                        region_id=region.id)
    return OrgCompanyService().create(short_name=f"Co DR-{suffix}",
                                      region_id=region.id,
                                      area_district_id=area.id)


@pytest.fixture()
def env(db):
    # The importer never sends a photo/license, same as the established
    # bulk-import convention (tests/unit/test_api_driver_import.py).
    _param(db, "REQUIRE_ASSIGNEE_PHOTO", "NO")
    _param(db, "REQUIRE_DRIVER_LICENSE", "NO")
    company_a = _company("A")
    company_b = _company("B")
    branch_in_a = BranchService().create(code="BR-DRV-A", name="In Company A",
                                         company_id=company_a.id)
    branch_in_b = BranchService().create(code="BR-DRV-B", name="In Company B",
                                         company_id=company_b.id)
    branch_unmapped = BranchService().create(code="BR-DRV-U",
                                             name="Unmapped Branch")
    db.session.commit()
    return {"company_a": company_a, "company_b": company_b,
            "branch_in_a": branch_in_a, "branch_in_b": branch_in_b,
            "branch_unmapped": branch_unmapped}


def _sheet(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Drivers"
    headers = [c for c, _req, _note in TEMPLATE_COLUMNS]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _row(employee_number, branch_code):
    # assignee_type EMPLOYEE (not DRIVER) so these company/branch-scoping
    # tests aren't entangled with the separate license-number requirement
    # that only applies to DRIVER-type assignees.
    return {"employee_number": employee_number, "first_name": "Juan",
            "last_name": "Dela Cruz", "branch_code": branch_code,
            "assignee_type": "EMPLOYEE", "status": "ACTIVE"}


# ── import_drivers: company_id omitted -- unchanged behaviour ────────

def test_without_a_company_id_an_unmapped_branch_still_imports(app, db, env):
    buf = _sheet([_row("DRV-U1", "BR-DRV-U")])
    stats = import_drivers(buf, dry_run=False)
    assert stats["created"] == 1
    assert stats["errors"] == []


# ── import_drivers: company_id given -- strict Company scoping ──────

def test_with_a_company_id_a_branch_in_that_company_imports(app, db, env):
    buf = _sheet([_row("DRV-A1", "BR-DRV-A")])
    stats = import_drivers(buf, dry_run=False, company_id=env["company_a"].id)
    assert stats["created"] == 1
    assert stats["errors"] == []


def test_with_a_company_id_a_branch_in_a_different_company_is_rejected(
        app, db, env):
    buf = _sheet([_row("DRV-A2", "BR-DRV-B")])
    stats = import_drivers(buf, dry_run=False, company_id=env["company_a"].id)
    assert stats["created"] == 0
    assert stats["skipped"] == 1
    assert "does not belong to the selected Company" in \
        stats["errors"][0]["problems"][0]


def test_with_a_company_id_an_unmapped_branch_is_rejected(app, db, env):
    buf = _sheet([_row("DRV-A3", "BR-DRV-U")])
    stats = import_drivers(buf, dry_run=False, company_id=env["company_a"].id)
    assert stats["created"] == 0
    assert stats["skipped"] == 1
    assert "does not belong to the selected Company" in \
        stats["errors"][0]["problems"][0]


def test_with_a_company_id_a_genuinely_unknown_branch_code_still_says_unknown(
        app, db, env):
    buf = _sheet([_row("DRV-A4", "NOPE-NOPE")])
    stats = import_drivers(buf, dry_run=False, company_id=env["company_a"].id)
    assert stats["created"] == 0
    assert "unknown branch code" in stats["errors"][0]["problems"][0]


# ── build_template: Reference sheet scoped to the selected Company ──

def test_template_without_a_company_id_lists_every_branch(app, db, env):
    wb = load_workbook(BytesIO(build_template()))
    refs = " ".join(str(c.value) for row in wb["Reference"].iter_rows()
                    for c in row if c.value)
    assert "BR-DRV-A" in refs
    assert "BR-DRV-B" in refs
    assert "BR-DRV-U" in refs


def test_template_with_a_company_id_lists_only_that_companys_branches(
        app, db, env):
    wb = load_workbook(BytesIO(build_template(company_id=env["company_a"].id)))
    refs = " ".join(str(c.value) for row in wb["Reference"].iter_rows()
                    for c in row if c.value)
    assert "BR-DRV-A" in refs
    assert "BR-DRV-B" not in refs
    assert "BR-DRV-U" not in refs
