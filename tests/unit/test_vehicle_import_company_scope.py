"""Batch 9 (2026-10-31): Vehicle CSV/Excel import scoped to a selected
Company.

Michael: "CSV import that supposed to be based on the selected
company" -- confirmed design: `company_id` is an OPTIONAL parameter to
import_vehicles()/build_template(). Omitted, behaviour is unchanged
(backward compatible with every existing caller/test). Given, every
branch_code in the file must resolve to a Branch under THAT Company --
a branch that exists but belongs to a different Company, or to none
at all (an "unmapped" branch), is rejected as a row error rather than
silently accepted, closing the anchorless-import gap flagged earlier
in this engagement.
"""
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.master_data.org.service import (
    BranchService, RegionService, AreaDistrictService, OrgCompanyService)
from app.modules.master_data.reference.service import VehicleTypeService
from app.modules.master_data.vehicle.import_service import (
    TEMPLATE_COLUMNS, build_template, import_vehicles)


def _company(suffix):
    region = RegionService().create(code=f"RG-VIC{suffix}",
                                     name=f"Region VIC {suffix}")
    area = AreaDistrictService().create(code=f"AD-VIC{suffix}",
                                        name=f"Area VIC {suffix}",
                                        region_id=region.id)
    return OrgCompanyService().create(short_name=f"Co VI-{suffix}",
                                      region_id=region.id,
                                      area_district_id=area.id)


@pytest.fixture()
def env(db):
    company_a = _company("A")
    company_b = _company("B")
    branch_in_a = BranchService().create(code="BR-VIC-A", name="In Company A",
                                         company_id=company_a.id)
    branch_in_b = BranchService().create(code="BR-VIC-B", name="In Company B",
                                         company_id=company_b.id)
    branch_unmapped = BranchService().create(code="BR-VIC-U",
                                             name="Unmapped Branch")
    vt = VehicleTypeService().create(code="LV-VIC", name="Light",
                                     category="LIGHT")
    db.session.commit()
    return {"company_a": company_a, "company_b": company_b,
            "branch_in_a": branch_in_a, "branch_in_b": branch_in_b,
            "branch_unmapped": branch_unmapped, "vt": vt}


def _sheet(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Vehicles"
    headers = [c for c, _req, _note in TEMPLATE_COLUMNS]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _row(conduction, branch_code, vt_code):
    return {"conduction_number": conduction, "vehicle_type_code": vt_code,
            "branch_code": branch_code, "brand": "Toyota", "model": "Vios",
            "year": 2022}


# ── import_vehicles: company_id omitted -- unchanged behaviour ───────

def test_without_a_company_id_an_unmapped_branch_still_imports(app, db, env):
    buf = _sheet([_row("VIC-U1", "BR-VIC-U", "LV-VIC")])
    stats = import_vehicles(buf, dry_run=False)
    assert stats["created"] == 1
    assert stats["errors"] == []


# ── import_vehicles: company_id given -- strict Company scoping ─────

def test_with_a_company_id_a_branch_in_that_company_imports(app, db, env):
    buf = _sheet([_row("VIC-A1", "BR-VIC-A", "LV-VIC")])
    stats = import_vehicles(buf, dry_run=False, company_id=env["company_a"].id)
    assert stats["created"] == 1
    assert stats["errors"] == []


def test_with_a_company_id_a_branch_in_a_different_company_is_rejected(
        app, db, env):
    buf = _sheet([_row("VIC-A2", "BR-VIC-B", "LV-VIC")])
    stats = import_vehicles(buf, dry_run=False, company_id=env["company_a"].id)
    assert stats["created"] == 0
    assert stats["skipped"] == 1
    assert "does not belong to the selected Company" in \
        stats["errors"][0]["problems"][0]


def test_with_a_company_id_an_unmapped_branch_is_rejected(app, db, env):
    buf = _sheet([_row("VIC-A3", "BR-VIC-U", "LV-VIC")])
    stats = import_vehicles(buf, dry_run=False, company_id=env["company_a"].id)
    assert stats["created"] == 0
    assert stats["skipped"] == 1
    assert "does not belong to the selected Company" in \
        stats["errors"][0]["problems"][0]


def test_with_a_company_id_a_genuinely_unknown_branch_code_still_says_unknown(
        app, db, env):
    buf = _sheet([_row("VIC-A4", "NOPE-NOPE", "LV-VIC")])
    stats = import_vehicles(buf, dry_run=False, company_id=env["company_a"].id)
    assert stats["created"] == 0
    assert "unknown branch_code" in stats["errors"][0]["problems"][0]


# ── build_template: Reference sheet scoped to the selected Company ──

def test_template_without_a_company_id_lists_every_branch(app, db, env):
    wb = load_workbook(BytesIO(build_template()))
    refs = " ".join(str(c.value) for row in wb["Reference"].iter_rows()
                    for c in row if c.value)
    assert "BR-VIC-A" in refs
    assert "BR-VIC-B" in refs
    assert "BR-VIC-U" in refs


def test_template_with_a_company_id_lists_only_that_companys_branches(
        app, db, env):
    wb = load_workbook(BytesIO(build_template(company_id=env["company_a"].id)))
    refs = " ".join(str(c.value) for row in wb["Reference"].iter_rows()
                    for c in row if c.value)
    assert "BR-VIC-A" in refs
    assert "BR-VIC-B" not in refs
    assert "BR-VIC-U" not in refs
