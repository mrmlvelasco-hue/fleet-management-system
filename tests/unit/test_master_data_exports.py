"""Master data configuration exports.

Requested so someone reviewing how the system is set up -- typically
the client checking PM scope configuration -- can extract what's
actually in the system rather than paging through a screen.
"""
import io

import openpyxl
import pytest

from app.core.security.registry import sync_permissions
from app.cli import _seed_admin
from app.core.reporting.master_data_exports import MASTER_DATA_EXPORTS


def _client(app, db):
    sync_permissions(); db.session.commit()
    _seed_admin("Testpass123!")
    c = app.test_client()
    c.post("/login", data={"username": "admin", "password": "Testpass123!"},
          follow_redirects=True)
    return c


@pytest.fixture()
def scope_data(app, db):
    from app.modules.master_data.reference.service import MaintenanceTypeService
    from app.modules.maintenance_config.service import PMScopeTemplateService
    mt = MaintenanceTypeService().create(code="PM-EXP", name="PM Export",
                                         category="PREVENTIVE")
    PMScopeTemplateService().create(
        maintenance_type_id=mt.id, name="Toyota Vios 5,000 km",
        items=[{"activity_code": f"A-{i:03d}",
                "activity_description": f"Check item {i}",
                "sort_order": i} for i in range(5)])
    return mt


@pytest.mark.parametrize("key", sorted(MASTER_DATA_EXPORTS))
def test_every_registered_export_produces_a_valid_workbook(app, db, key):
    """Every module in the registry must actually export. A wrong field
    name would only surface here, not at import time."""
    r = _client(app, db).get(f"/master/exports/{key}.xlsx")
    assert r.status_code == 200, f"{key} returned {r.status_code}"
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb.worksheets, f"{key} produced no sheets"


def test_pm_scope_export_includes_the_activity_lines(app, db, scope_data):
    """The point of this export: a list of template NAMES would not
    answer "what does this PM actually cover" -- the activities are the
    configuration."""
    r = _client(app, db).get("/master/exports/pm-scope-templates.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert "Scope Activities" in wb.sheetnames
    ws = wb["Scope Activities"]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    assert len(rows) == 5, "activity lines missing from the export"
    descriptions = [r[4] for r in rows]
    assert "Check item 0" in descriptions


def test_pm_scope_export_activity_rows_name_their_template(app, db,
                                                           scope_data):
    """Each activity row has to be traceable to its template, or the
    second sheet is an unattributable list of tasks."""
    r = _client(app, db).get("/master/exports/pm-scope-templates.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb["Scope Activities"]
    for row in ws.iter_rows(min_row=4, values_only=True):
        assert row[0] == "Toyota Vios 5,000 km"


def test_export_headers_are_present_and_frozen(app, db, scope_data):
    r = _client(app, db).get("/master/exports/pm-scope-templates.xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(r.data))["Scope Templates"]
    headers = [c.value for c in ws[3]]
    assert "Template Name" in headers
    assert "Activities" in headers
    assert ws.freeze_panes == "A4", "headers should stay visible when scrolling"


def test_unknown_export_key_is_a_404_not_an_empty_file(app, db):
    """An empty spreadsheet would read as "nothing is configured",
    which is a materially different and wrong message."""
    r = _client(app, db).get("/master/exports/not-a-module.xlsx")
    assert r.status_code == 404


def test_export_requires_the_matching_view_permission(app, db):
    """An export is a full dump of a module, so it must need the same
    permission as viewing that module -- not a blanket export right."""
    from app.modules.user_management.models import User
    from app.core.security.password import hash_password
    sync_permissions(); db.session.commit()
    user = User(username="noview", email="nv@x.com",
                password_hash=hash_password("Testpass123!"),
                must_change_password=False)
    db.session.add(user)
    db.session.commit()
    c = app.test_client()
    c.post("/login", data={"username": "noview",
                           "password": "Testpass123!"})
    r = c.get("/master/exports/vehicles.xlsx")
    assert r.status_code == 403


def test_a_missing_related_record_does_not_break_the_export(app, db):
    """A template with no linked PM schedule must still export, showing
    a dash -- one incomplete record shouldn't cost the whole file."""
    from app.modules.master_data.reference.service import MaintenanceTypeService
    from app.modules.maintenance_config.service import PMScopeTemplateService
    mt = MaintenanceTypeService().create(code="PM-ORPH", name="Orphan",
                                         category="PREVENTIVE")
    PMScopeTemplateService().create(
        maintenance_type_id=mt.id, name="Unlinked Template",
        items=[{"activity_code": "A-1", "activity_description": "Work",
                "sort_order": 0}])
    r = _client(app, db).get("/master/exports/pm-scope-templates.xlsx")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.data))["Scope Templates"]
    names = [row[0] for row in ws.iter_rows(min_row=4, values_only=True)]
    assert "Unlinked Template" in names


def test_export_button_appears_on_the_list_pages(app, db, scope_data):
    client = _client(app, db)
    for url, key in [("/admin/pm-scope-templates", "pm-scope-templates"),
                     ("/master/vehicles", "vehicles"),
                     ("/master/departments", "departments")]:
        html = client.get(url).get_data(as_text=True)
        assert f"/master/exports/{key}.xlsx" in html, f"no export button on {url}"
