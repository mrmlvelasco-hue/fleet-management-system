"""Filter-aware Excel export of the Vehicle list.

Flask's Export button (`master_data_export(key='vehicles')`) dumps the
whole vehicles table. It sits directly beneath three filter dropdowns and
ignores all of them, which makes the behaviour actively misleading:
exporting 5,000 rows while the user is looking at 12 is the failure that
gets noticed in front of a client.

React diverges deliberately. Export respects the active filter, and the
unfiltered export stays reachable by clearing the filters -- so nothing
Flask could produce becomes unreachable.

The endpoint takes the SAME parameters as GET /api/v1/vehicles and calls
the SAME service method. A second query shaped slightly differently would
mean the spreadsheet and the screen could disagree about what the filter
means, and the spreadsheet is the copy that gets emailed on.
"""
import json
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.security.password import hash_password
from app.core.security.registry import sync_permissions
from app.modules.master_data.org.service import BranchService
from app.modules.master_data.reference.service import VehicleTypeService
from app.modules.master_data.vehicle.service import VehicleService
from app.modules.user_management.models import Permission, Role, User


@pytest.fixture()
def exp_env(db):
    sync_permissions()
    db.session.commit()

    role = Role(name="Vehicle Exporter")
    role.permissions = Permission.query.filter(
        Permission.code == "vehicle.view").all()
    user = User(username="exporter", email="exp@e.com",
                password_hash=hash_password("secret123"), is_active=True)
    user.roles = [role]

    nobody_role = Role(name="No Access At All")
    nobody = User(username="noexp", email="nx@e.com",
                  password_hash=hash_password("secret123"), is_active=True)
    nobody.roles = [nobody_role]
    db.session.add_all([role, user, nobody_role, nobody])

    light = VehicleTypeService().create(code="LV-EX", name="Light",
                                        category="LIGHT")
    heavy = VehicleTypeService().create(code="HV-EX", name="Heavy",
                                        category="HEAVY")
    manila = BranchService().create(code="BR-MNL", name="Manila")
    cebu = BranchService().create(code="BR-CEB", name="Cebu")

    VehicleService().create(vehicle_type_id=light.id, brand="Toyota",
                            model="Hilux", year=2021, branch_id=manila.id,
                            plate_number="AAA-1111", status="ACTIVE")
    VehicleService().create(vehicle_type_id=heavy.id, brand="Isuzu",
                            model="Forward", year=2019, branch_id=cebu.id,
                            plate_number="BBB-2222", status="IN_REPAIR")
    VehicleService().create(vehicle_type_id=light.id, brand="Toyota",
                            model="Vios", year=2018, branch_id=manila.id,
                            plate_number="CCC-3333", status="DISPOSED")
    db.session.commit()
    return {"user": user, "manila": manila, "cebu": cebu,
            "light": light, "heavy": heavy}


def _token(client, username="exporter"):
    r = client.post("/api/v1/auth/token",
                    json={"username": username, "password": "secret123"})
    return json.loads(r.get_data(as_text=True)).get("access_token")


def _plates(response):
    """Plate column of the returned workbook."""
    wb = load_workbook(BytesIO(response.data))
    ws = wb.active
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        values = [c.value for c in row]
        if "Plate No." in values:
            header_row = row[0].row
            column = values.index("Plate No.") + 1
            break
    assert header_row, "no header row found in the exported workbook"
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(r, column).value
        if value:
            out.append(value)
    return out


def _export(client, token, query=""):
    return client.get(f"/api/v1/vehicles/export.xlsx{query}",
                      headers={"Authorization": f"Bearer {token}"})


# ── Auth ────────────────────────────────────────────────────────────────────

def test_export_rejects_anonymous(db, client, exp_env):
    assert client.get("/api/v1/vehicles/export.xlsx").status_code == 401


def test_export_requires_vehicle_view(db, client, exp_env):
    assert _export(client, _token(client, "noexp")).status_code == 403


# ── Content type ────────────────────────────────────────────────────────────

def test_export_returns_a_spreadsheet_attachment(db, client, exp_env):
    r = _export(client, _token(client))
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["Content-Type"]
    assert "attachment" in r.headers["Content-Disposition"]
    assert ".xlsx" in r.headers["Content-Disposition"]


# ── The point of the endpoint ───────────────────────────────────────────────

def test_unfiltered_export_holds_the_active_fleet(db, client, exp_env):
    """No filters means what the list means with no filters: disposed
    vehicles stay out until asked for."""
    plates = _plates(_export(client, _token(client)))
    assert set(plates) == {"AAA-1111", "BBB-2222"}


def test_export_respects_the_branch_filter(db, client, exp_env):
    plates = _plates(_export(client, _token(client),
                             f"?branch_id={exp_env['manila'].id}"))
    assert plates == ["AAA-1111"]


def test_export_respects_the_status_filter(db, client, exp_env):
    plates = _plates(_export(client, _token(client), "?status=IN_REPAIR"))
    assert plates == ["BBB-2222"]


def test_export_respects_the_vehicle_type_filter(db, client, exp_env):
    plates = _plates(_export(client, _token(client),
                             f"?vehicle_type_id={exp_env['heavy'].id}"))
    assert plates == ["BBB-2222"]


def test_export_respects_the_search_term(db, client, exp_env):
    plates = _plates(_export(client, _token(client), "?q=Hilux"))
    assert plates == ["AAA-1111"]


def test_export_respects_show_disposed(db, client, exp_env):
    """The disposed toggle is explicit on the list and must be explicit
    here. A register that quietly includes retired vehicles is a
    different document from the one on screen."""
    plates = _plates(_export(client, _token(client), "?show_disposed=1"))
    assert "CCC-3333" in plates


def test_export_combines_filters(db, client, exp_env):
    """Filters narrow together, as they do on the list. If they were
    OR-ed, an export would grow when the user narrowed the screen."""
    plates = _plates(_export(
        client, _token(client),
        f"?branch_id={exp_env['manila'].id}&status=IN_REPAIR"))
    assert plates == []


def test_export_is_not_capped_at_a_page(db, client, exp_env):
    """The list pages at 20 and caps page_size at 200. An export that
    inherited either would hand over a truncated file that looks
    complete -- the worst possible failure for a document someone
    forwards."""
    light = exp_env["light"]
    manila = exp_env["manila"]
    for i in range(250):
        VehicleService().create(vehicle_type_id=light.id, brand="Toyota",
                                model="Hiace", year=2020,
                                branch_id=manila.id,
                                plate_number=f"BULK-{i:04d}")
    db.session.commit()
    plates = _plates(_export(client, _token(client), "?q=Hiace"))
    assert len(plates) == 250


def test_export_rejects_an_unknown_status(db, client, exp_env):
    """Same rejection the list gives. Silently ignoring a typo would
    export the whole fleet under a filter the caller believes applied."""
    assert _export(client, _token(client), "?status=NONSENSE").status_code == 400


def test_export_rejects_a_non_integer_branch(db, client, exp_env):
    assert _export(client, _token(client), "?branch_id=abc").status_code == 400


def test_export_cannot_widen_org_scope(db, client, exp_env):
    """Scoping comes from VehicleService.list_page(user=...), the same
    call the list makes. An export must never be a way to obtain rows
    the screen would not show.

    The out-of-scope vehicle is given a plate that sorts FIRST. An
    earlier version of this test let the in-scope vehicle sort first,
    and passed even with scoping removed from the fetch: the row count
    was still correct, so slicing the unscoped set simply returned the
    right row by accident. Caught by mutation check.
    """
    from app.modules.user_management.org_scope_service import (
        UserOrgScopeService)

    VehicleService().create(vehicle_type_id=exp_env["light"].id,
                            brand="Isuzu", model="Crosswind", year=2020,
                            branch_id=exp_env["cebu"].id,
                            plate_number="AAA-0001")
    VehicleService().create(vehicle_type_id=exp_env["light"].id,
                            brand="Toyota", model="Fortuner", year=2022,
                            branch_id=exp_env["manila"].id,
                            plate_number="ZZZ-9999")
    UserOrgScopeService().assign(exp_env["user"].id, scope_type="BRANCH",
                                 branch_id=exp_env["manila"].id)
    db.session.commit()

    plates = _plates(_export(client, _token(client)))
    assert "ZZZ-9999" in plates, "the caller's own branch went missing"
    assert "AAA-0001" not in plates
    assert "BBB-2222" not in plates


# ── Disposed count for the list toggle ──────────────────────────────────────

def test_summary_carries_a_disposed_count(db, client, exp_env):
    """The Jinja list badges its Show/Hide Disposed toggle with a count,
    and says why in a comment: without it, clicking when there happen to
    be zero disposed vehicles looks exactly like the click did nothing,
    which is indistinguishable from a broken filter.

    Flask computes it by listing the whole fleet a SECOND time. Folded
    into the summary here instead -- that endpoint already exists, is
    already off the critical path, and a third full list per page render
    is not worth a badge.
    """
    r = client.get("/api/v1/vehicles/summary",
                   headers={"Authorization": f"Bearer {_token(client)}"})
    assert r.status_code == 200
    assert json.loads(r.get_data(as_text=True))["disposed"] == 1


def test_disposed_count_ignores_the_status_filter(db, client, exp_env):
    """It answers "how many are there to reveal", not "how many match
    the current filter". Tying it to the filter would show 0 whenever
    any status was selected, and the toggle would look broken in exactly
    the situation it exists for."""
    r = client.get("/api/v1/vehicles/summary?status=ACTIVE",
                   headers={"Authorization": f"Bearer {_token(client)}"})
    assert json.loads(r.get_data(as_text=True))["disposed"] == 1


def test_disposed_count_respects_org_scope(db, client, exp_env):
    """A count is still a disclosure: it must not reveal that another
    branch has retired vehicles."""
    from app.modules.user_management.org_scope_service import (
        UserOrgScopeService)
    UserOrgScopeService().assign(exp_env["user"].id, scope_type="BRANCH",
                                 branch_id=exp_env["cebu"].id)
    db.session.commit()
    r = client.get("/api/v1/vehicles/summary",
                   headers={"Authorization": f"Bearer {_token(client)}"})
    # The only disposed vehicle is in Manila.
    assert json.loads(r.get_data(as_text=True))["disposed"] == 0
