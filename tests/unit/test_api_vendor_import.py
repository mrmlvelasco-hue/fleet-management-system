"""Vendor bulk import from Excel (B24).

Same shape as the vehicle import (tests/unit/test_api_vehicle_import.py),
because the client asked for "the same UI as the vehicle upload":

  GET  /api/v1/vendors/import/template  -> fill-in .xlsx, with this
       install's own Branch codes on a "Reference" sheet
  POST /api/v1/vendors/import           -> multipart `file`; a PREVIEW
       unless commit=1, so uploading the wrong file saves nothing

Agreed with the client: a code that already exists is SKIPPED and
reported -- an import never changes an existing vendor. A row that fails
validation is skipped; the rest still import (the vehicle importer's
partial-import rule). Gated on vendor.create, as creating one vendor is.
"""
import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.core.security.password import hash_password
from app.core.security.registry import sync_permissions
from app.modules.master_data.org.service import BranchService
from app.modules.master_data.vendor.import_service import TEMPLATE_COLUMNS
from app.modules.master_data.vendor.models import Vendor
from app.modules.master_data.vendor.service import VendorService
from app.modules.user_management.models import Permission, Role, User

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def env(db):
    sync_permissions()
    BranchService().create(code="BR-MUN", name="Muntinlupa")
    BranchService().create(code="BR-LP", name="Las Piñas")
    VendorService().create(code="VEN-OLD", name="Already Here Co",
                           contact_person="Keep Me")

    def mk(username, codes):
        role = Role(name=f"role-{username}")
        role.permissions = Permission.query.filter(
            Permission.code.in_(codes)).all()
        u = User(username=username, email=f"{username}@e.com",
                 password_hash=hash_password("secret123"), is_active=True)
        u.roles = [role]
        db.session.add_all([role, u])

    mk("vimporter", ["vendor.view", "vendor.create"])
    mk("vviewer", ["vendor.view"])
    db.session.commit()


def _hdr(client, username="vimporter"):
    r = client.post("/api/v1/auth/token",
                    json={"username": username, "password": "secret123"})
    tok = json.loads(r.get_data(as_text=True))["access_token"]
    return {"Authorization": f"Bearer {tok}"}


def _sheet(rows, headers=None, title="Vendors"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    headers = headers or [c for c, _req, _note in TEMPLATE_COLUMNS]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


GOOD = {"code": "VEN-101", "name": "Alabang Toyota", "vendor_type": "services",
        "tin": "123-456-789-000", "contact_person": "Juan Dela Cruz",
        "phone": "(02) 8807-1234", "email": "service@alabangtoyota.ph",
        "address": "Alabang-Zapote Rd", "city": "Muntinlupa",
        "branch_codes": "br-mun, BR-LP"}


def _post(client, buf, commit=None, user="vimporter", name="vendors.xlsx"):
    data = {"file": (buf, name)}
    if commit is not None:
        data["commit"] = commit
    return client.post("/api/v1/vendors/import", headers=_hdr(client, user),
                       data=data, content_type="multipart/form-data")


def _problems(body):
    return {e["identifier"]: " | ".join(e["problems"]) for e in body["errors"]}


# ── template ────────────────────────────────────────────────────────────

def test_template_is_a_workbook_with_this_installs_branch_codes(client, db, env):
    r = client.get("/api/v1/vendors/import/template", headers=_hdr(client))
    assert r.status_code == 200
    assert XLSX in r.headers["Content-Type"]
    assert "Vendor_Import_Template.xlsx" in r.headers["Content-Disposition"]
    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames[:2] == ["Vendors", "Reference"]
    headers = [c.value for c in wb["Vendors"][1]]
    assert headers == [c for c, _r, _n in TEMPLATE_COLUMNS]
    refs = " ".join(str(c.value) for row in wb["Reference"].iter_rows()
                    for c in row if c.value)
    assert "BR-MUN" in refs and "BR-LP" in refs
    assert "GOODS" in refs and "SERVICES" in refs and "BOTH" in refs


def test_template_requires_vendor_create(client, db, env):
    r = client.get("/api/v1/vendors/import/template",
                   headers=_hdr(client, "vviewer"))
    assert r.status_code == 403


def test_the_templates_own_example_row_is_never_imported(client, db, env):
    """People forget to delete it -- uploading the untouched template
    must import nothing, not a vendor called EXAMPLE."""
    tmpl = client.get("/api/v1/vendors/import/template", headers=_hdr(client)).data
    r = _post(client, BytesIO(tmpl), commit="1")
    body = r.get_json()
    assert body["total_rows"] == 0 and body["created"] == 0
    assert Vendor.query.filter(Vendor.code.like("EXAMPLE%")).count() == 0


# ── preview vs commit ───────────────────────────────────────────────────

def test_preview_is_the_default_and_saves_nothing(client, db, env):
    r = _post(client, _sheet([GOOD]))
    assert r.status_code == 200, r.get_data(as_text=True)
    body = r.get_json()
    assert body["dry_run"] is True
    assert body["created"] == 1 and body["skipped"] == 0
    assert Vendor.query.filter_by(code="VEN-101").count() == 0


def test_commit_zero_is_still_a_preview(client, db, env):
    _post(client, _sheet([GOOD]), commit="0")
    assert Vendor.query.filter_by(code="VEN-101").count() == 0


def test_commit_saves_every_field_and_the_branches(client, db, env):
    r = _post(client, _sheet([GOOD]), commit="1")
    body = r.get_json()
    assert body["dry_run"] is False and body["created"] == 1
    v = Vendor.query.filter_by(code="VEN-101").one()
    assert v.name == "Alabang Toyota"
    assert v.vendor_type == "SERVICES"          # normalised to upper case
    assert (v.tin, v.contact_person, v.phone, v.email, v.address, v.city) == (
        "123-456-789-000", "Juan Dela Cruz", "(02) 8807-1234",
        "service@alabangtoyota.ph", "Alabang-Zapote Rd", "Muntinlupa")
    assert sorted(b.code for b in v.branches) == ["BR-LP", "BR-MUN"]


def test_blank_type_defaults_to_goods_and_blank_branches_serve_everyone(
        client, db, env):
    _post(client, _sheet([{"code": "VEN-102", "name": "Plain Supplier"}]),
          commit="1")
    v = Vendor.query.filter_by(code="VEN-102").one()
    assert v.vendor_type == "GOODS"
    assert v.branches == []


# ── row validation: skipped and reported, the rest still import ─────────

def test_an_existing_code_is_skipped_and_the_existing_vendor_untouched(
        client, db, env):
    r = _post(client, _sheet([{"code": "ven-old", "name": "Overwrite Attempt",
                               "contact_person": "Someone Else"}, GOOD]),
              commit="1")
    body = r.get_json()
    assert body["created"] == 1 and body["skipped"] == 1
    # Reported as written in the file, so it can be found in the sheet;
    # matched case-insensitively, so "ven-old" is still VEN-OLD.
    assert "already exists" in _problems(body)["ven-old"]
    old = Vendor.query.filter_by(code="VEN-OLD").one()
    assert (old.name, old.contact_person) == ("Already Here Co", "Keep Me")


def test_a_code_repeated_in_the_same_file_imports_once(client, db, env):
    r = _post(client, _sheet([GOOD, {**GOOD, "name": "Second Copy"}]),
              commit="1")
    body = r.get_json()
    assert body["created"] == 1 and body["skipped"] == 1
    assert "appears more than once" in _problems(body)["VEN-101"]
    assert Vendor.query.filter_by(code="VEN-101").one().name == "Alabang Toyota"


def test_each_bad_row_names_its_own_problems(client, db, env):
    rows = [
        {"code": "VEN-201"},                                       # no name
        {"name": "No Code Inc"},                                   # no code
        {"code": "VEN-202", "name": "X", "vendor_type": "PARTS"},  # bad type
        {"code": "VEN-203", "name": "X", "branch_codes": "BR-MUN, NOWHERE"},
        {"code": "VEN-204", "name": "X", "email": "not-an-email"},
        {"code": "V" * 21, "name": "X"},                           # > 20 chars
        GOOD,
    ]
    body = _post(client, _sheet(rows), commit="1").get_json()
    assert body["total_rows"] == 7
    assert body["created"] == 1 and body["skipped"] == 6
    p = _problems(body)
    assert "name is required" in p["VEN-201"]
    assert "code is required" in p["No Code Inc"]
    assert "vendor_type 'PARTS'" in p["VEN-202"]
    assert "unknown branch code 'NOWHERE'" in p["VEN-203"]
    assert "email" in p["VEN-204"]
    assert "at most 20" in p["V" * 21]
    # Only the good row was saved -- and none of the bad ones, even partly.
    assert Vendor.query.filter(Vendor.code.like("VEN-2%")).count() == 0
    assert Vendor.query.filter_by(code="VEN-101").count() == 1


def test_preview_reports_the_same_problems_a_commit_would(client, db, env):
    rows = [GOOD, {"code": "VEN-202", "name": "X", "vendor_type": "PARTS"}]
    preview = _post(client, _sheet(rows)).get_json()
    assert (preview["created"], preview["skipped"]) == (1, 1)
    assert "vendor_type 'PARTS'" in _problems(preview)["VEN-202"]


def test_blank_rows_are_ignored(client, db, env):
    body = _post(client, _sheet([{}, GOOD, {}])).get_json()
    assert body["total_rows"] == 1


# ── file-level problems ─────────────────────────────────────────────────

def test_missing_required_columns_is_a_400_naming_them(client, db, env):
    r = _post(client, _sheet([{"code": "A"}], headers=["code", "city"]))
    assert r.status_code == 400
    assert "name" in r.get_json()["message"]


def test_a_file_that_is_not_a_workbook_is_a_400(client, db, env):
    r = _post(client, BytesIO(b"not a spreadsheet"), name="vendors.xlsx")
    assert r.status_code == 400
    assert "Could not read the file" in r.get_json()["message"]


def test_no_file_is_a_400(client, db, env):
    r = client.post("/api/v1/vendors/import", headers=_hdr(client),
                    data={}, content_type="multipart/form-data")
    assert r.status_code == 400


def test_import_requires_vendor_create(client, db, env):
    r = _post(client, _sheet([GOOD]), commit="1", user="vviewer")
    assert r.status_code == 403
    assert Vendor.query.filter_by(code="VEN-101").count() == 0


def test_a_first_sheet_with_another_name_is_still_read(client, db, env):
    """The vehicle importer's rule: 'Vendors' if present, else sheet 1."""
    body = _post(client, _sheet([GOOD], title="Sheet1")).get_json()
    assert body["created"] == 1
