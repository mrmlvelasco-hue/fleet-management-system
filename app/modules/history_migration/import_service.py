"""Import service for the Vehicle History Migration template.

Reads two sheets -- Maintenance History and Registration History -- and
creates real MaintenanceOrder / VehicleRegistration rows from them,
tagged as historical so they never enter the live approval workflow or
appear in "For My Action".

Design decisions worth stating plainly:

  * Historical rows are created directly as COMPLETED with no
    approval_instance. The work already happened, sometimes years ago
    -- routing it through Draft -> Submitted -> Approved would be
    fiction, and would pollute this period's real approval statistics.

  * document_number is left NULL for historical rows rather than
    consuming a real number from the live numbering counter. Running a
    bulk historical import through AutoNumberingService would either
    hand out today's sequence to a 2019 record (nonsensical) or require
    rolling the counter back per period (fragile, and risks exactly the
    kind of numbering collision already fixed elsewhere in this
    project). A historical row's identity is its vehicle + date, not a
    generated number.

  * Every row from one file shares one HistoricalImportBatch, so a
    wrong upload can be undone as a single action across BOTH
    Maintenance and Registration history at once -- the same guarantee
    already built for fuel statement imports, and for the same reason:
    a bulk historical migration is exactly the kind of one-shot,
    high-stakes operation where an undo matters most.

  * A row that fails validation is skipped and reported, not fatal to
    the whole file. On a few hundred hand-transcribed rows, one bad
    date must not block the other 299.
"""
from datetime import datetime, date
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.extensions import db


def _norm(value):
    return str(value).strip() if value is not None else ""


def _to_date(value, field_name):
    if value is None or _norm(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _norm(value)
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"{field_name}: '{text}' is not a recognisable date "
                    f"(use YYYY-MM-DD)")


def _to_int(value, field_name):
    if value is None or _norm(value) == "":
        return None
    try:
        return int(Decimal(str(value).replace(",", "")))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{field_name}: '{value}' is not a whole number")


def _to_decimal(value, field_name):
    if value is None or _norm(value) == "":
        return None
    try:
        return Decimal(str(value).replace(",", ""))
    except InvalidOperation:
        raise ValueError(f"{field_name}: '{value}' is not a valid amount")


def _find_vehicle(vehicles_by_key, plate_or_conduction):
    key = _norm(plate_or_conduction).upper()
    return vehicles_by_key.get(key)


def _load_vehicle_index():
    """{plate or conduction number, uppercased: Vehicle}."""
    from app.modules.master_data.vehicle.models import Vehicle
    index = {}
    for v in Vehicle.query.all():
        if v.plate_number:
            index[v.plate_number.strip().upper()] = v
        if v.conduction_number:
            index[v.conduction_number.strip().upper()] = v
    return index


def import_vehicle_history(file_stream, dry_run: bool = True,
                          filename: str = None, user_id: int = None) -> dict:
    """Import both sheets of the Vehicle History Migration template.

    Returns per-sheet stats and a combined error list. Nothing is
    written to the database when dry_run is True.
    """
    from app.modules.history_migration.models import HistoricalImportBatch
    from app.modules.transactions.maintenance_order.models import (
        MaintenanceOrder)
    from app.modules.transactions.vehicle_registration.models import (
        VehicleRegistration)

    wb = load_workbook(file_stream, data_only=True)
    vehicles = _load_vehicle_index()

    stats = {
        "maintenance": {"total_rows": 0, "created": 0, "skipped": 0},
        "registration": {"total_rows": 0, "created": 0, "skipped": 0},
        "errors": [],
    }

    batch = None
    if not dry_run:
        batch = HistoricalImportBatch(
            filename=filename, imported_by=user_id,
            imported_at=datetime.now())
        db.session.add(batch)
        db.session.flush()

    # ── Maintenance History ─────────────────────────────────────────
    if "Maintenance History" in wb.sheetnames:
        ws = wb["Maintenance History"]
        for row_number, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2):
            if row is None or all(_norm(v) == "" for v in row):
                continue
            if _norm(row[0]).startswith("^"):
                continue   # the shipped "these rows are examples" note
            (plate, service_date, mtype, description, odometer,
             cost, shop, reference) = (list(row) + [None] * 8)[:8]

            stats["maintenance"]["total_rows"] += 1
            errors = []

            vehicle = _find_vehicle(vehicles, plate)
            if vehicle is None:
                errors.append(f"no vehicle matches '{plate}'")

            try:
                parsed_date = _to_date(service_date, "Date of Service")
                if parsed_date is None:
                    errors.append("Date of Service is required")
            except ValueError as exc:
                errors.append(str(exc)); parsed_date = None

            mtype_norm = _norm(mtype).upper()
            if mtype_norm not in ("PREVENTIVE", "CORRECTIVE"):
                errors.append(
                    f"Maintenance Type must be PREVENTIVE or CORRECTIVE, "
                    f"got '{mtype}'")

            try:
                odometer_val = _to_int(odometer, "Odometer at Service")
            except ValueError as exc:
                errors.append(str(exc)); odometer_val = None
            try:
                cost_val = _to_decimal(cost, "Cost")
            except ValueError as exc:
                errors.append(str(exc)); cost_val = None

            if errors:
                stats["maintenance"]["skipped"] += 1
                stats["errors"].append({
                    "sheet": "Maintenance History", "row": row_number,
                    "identifier": plate or "(blank)", "problems": errors})
                continue

            if dry_run:
                stats["maintenance"]["created"] += 1
                continue

            order = MaintenanceOrder(
                vehicle_id=vehicle.id, order_category="MAINTENANCE",
                category=mtype_norm[:12],
                description=_norm(description) or None,
                odometer_at_service=odometer_val,
                scheduled_date=parsed_date, completed_date=parsed_date,
                assigned_mechanic=_norm(shop) or None,
                actual_cost=cost_val, status="COMPLETED",
                is_historical=True, historical_batch_id=batch.id)
            db.session.add(order)
            stats["maintenance"]["created"] += 1

    # ── Registration History ────────────────────────────────────────
    if "Registration History" in wb.sheetnames:
        ws = wb["Registration History"]
        for row_number, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2):
            if row is None or all(_norm(v) == "" for v in row):
                continue
            if _norm(row[0]).startswith("^"):
                continue   # the shipped "these rows are examples" note
            (plate, reg_type, reg_date, expiry, or_no, cr_no,
             cost) = (list(row) + [None] * 7)[:7]

            stats["registration"]["total_rows"] += 1
            errors = []

            vehicle = _find_vehicle(vehicles, plate)
            if vehicle is None:
                errors.append(f"no vehicle matches '{plate}'")

            reg_type_norm = _norm(reg_type).upper()
            if reg_type_norm not in ("NEW", "RENEWAL"):
                errors.append(
                    f"Registration Type must be NEW or RENEWAL, got "
                    f"'{reg_type}'")

            try:
                parsed_reg_date = _to_date(reg_date, "Registration Date")
                if parsed_reg_date is None:
                    errors.append("Registration Date is required")
            except ValueError as exc:
                errors.append(str(exc)); parsed_reg_date = None
            try:
                parsed_expiry = _to_date(expiry, "Expiry Date")
            except ValueError as exc:
                errors.append(str(exc)); parsed_expiry = None
            try:
                cost_val = _to_decimal(cost, "Cost")
            except ValueError as exc:
                errors.append(str(exc)); cost_val = None

            if errors:
                stats["registration"]["skipped"] += 1
                stats["errors"].append({
                    "sheet": "Registration History", "row": row_number,
                    "identifier": plate or "(blank)", "problems": errors})
                continue

            if dry_run:
                stats["registration"]["created"] += 1
                continue

            reg = VehicleRegistration(
                vehicle_id=vehicle.id, registration_type=reg_type_norm,
                registration_date=parsed_reg_date, expiry_date=parsed_expiry,
                or_number=_norm(or_no) or None,
                cr_number=_norm(cr_no) or None,
                or_cr_cost=cost_val, status="COMPLETED",
                is_historical=True, historical_batch_id=batch.id)
            db.session.add(reg)
            stats["registration"]["created"] += 1

    if not dry_run:
        batch.maintenance_rows = stats["maintenance"]["created"]
        batch.registration_rows = stats["registration"]["created"]
        db.session.commit()
        stats["batch_id"] = batch.id
    else:
        stats["batch_id"] = None

    return stats


def delete_historical_batch(batch_id):
    """Undo one historical import -- removes exactly and only the rows
    that one file created, across both Maintenance and Registration
    history."""
    from app.modules.history_migration.models import HistoricalImportBatch
    from app.modules.transactions.maintenance_order.models import (
        MaintenanceOrder)
    from app.modules.transactions.vehicle_registration.models import (
        VehicleRegistration)

    mo_rows = MaintenanceOrder.query.filter_by(
        historical_batch_id=batch_id).all()
    vr_rows = VehicleRegistration.query.filter_by(
        historical_batch_id=batch_id).all()
    count = len(mo_rows) + len(vr_rows)
    for row in mo_rows + vr_rows:
        db.session.delete(row)
    batch = db.session.get(HistoricalImportBatch, batch_id)
    if batch:
        db.session.delete(batch)
    db.session.commit()
    return count
