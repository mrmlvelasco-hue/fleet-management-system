"""Fuel transaction import and export.

Deliberately provider-agnostic. The client's Petron statement format
isn't confirmed yet, so this reads by COLUMN HEADER through an alias
table rather than by fixed position. Adding Petron's exact headers later
is a few entries in COLUMN_ALIASES, not a rewrite — and a statement from
Shell or Caltex would work the same way.

Everything imported still passes through OdometerValidationService, so a
mistyped pump reading is quarantined on the way in rather than being
discovered later in a consumption report.
"""
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from sqlalchemy import func

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from bs4 import BeautifulSoup

from app.extensions import db


# Header aliases, lower-cased and stripped. Providers name the same
# column half a dozen ways; matching on a set of known spellings means a
# new statement layout usually needs no code change at all.
COLUMN_ALIASES = {
    "transaction_date": ("transaction_date", "date", "txn date",
                        "transaction date", "posting date", "trans date",
                        "datetime", "date/time"),
    "card_number": ("card_number", "card no", "card no.", "card number",
                   "fleet card", "cardno", "card"),
    "plate_number": ("plate_number", "plate", "plate no", "plate no.",
                    "vehicle", "vehicle no", "conduction", "unit"),
    "station": ("station", "site", "site name", "station name",
               "merchant", "location", "branch"),
    "litres": ("litres", "liters", "volume", "qty", "quantity",
              "liters sold", "volume l"),
    "price_per_litre": ("price_per_litre", "unit price", "price l",
                       "price per liter", "price", "unit cost"),
    "total_amount": ("total_amount", "amount", "total", "gross amount",
                    "net amount", "total cost"),
    "odometer_reported": ("odometer_reported", "odometer", "odo",
                         "mileage", "km reading", "odometer reading",
                         "odometer reported"),
    "fuel_type": ("fuel_type", "product", "fuel", "product name",
                 "description", "fuel type"),
    "reference_number": ("reference_number", "reference", "ref no",
                        "ref no.", "receipt", "receipt no", "invoice no",
                        "transaction no", "or no", "trans slip no"),
    "driver_name": ("driver_name", "driver", "driver name", "cardholder"),
}

# Petron's own fuel-card statement export uses a provider-specific set of
# headers that don't fit the generic alias lists above cleanly -- added
# individually, verified against the client's real sample file, per the
# Fuel + Org master prompt's column-mapping spec:
#   TxnDate -> transaction_date, EmbName -> plate_number,
#   BillingTxnAmt -> total_amount, Descp -> fuel_type.
# VehRegsNo/VehRegsNo1, AcctNo, Charging Department and Station Address
# are deliberately NOT mapped -- no verified FMS field fits them, and
# EmbName (not VehRegsNo) is the primary plate source per the mapping
# spec. DriverCd is also deliberately unmapped: its values are "0" or a
# card's own last four digits, not a driver's name, and no verified
# code->driver lookup exists.
COLUMN_ALIASES["transaction_date"] = COLUMN_ALIASES["transaction_date"] + ("txndate",)
COLUMN_ALIASES["plate_number"] = COLUMN_ALIASES["plate_number"] + ("embname",)
COLUMN_ALIASES["total_amount"] = COLUMN_ALIASES["total_amount"] + ("billingtxnamt",)
COLUMN_ALIASES["fuel_type"] = COLUMN_ALIASES["fuel_type"] + ("descp",)

TEMPLATE_COLUMNS = [
    ("transaction_date", True, "YYYY-MM-DD or YYYY-MM-DD HH:MM"),
    ("plate_number", True, "Plate or Conduction number — must match a vehicle."),
    ("card_number", False, "Fleet card number, if the fill was on a card."),
    ("station", False, "Station or site name."),
    ("fuel_type", False, "e.g. DIESEL, GASOLINE"),
    ("litres", True, "Numbers only, e.g. 45.5"),
    ("price_per_litre", False, "Numbers only, e.g. 62.15"),
    ("total_amount", True, "Numbers only, e.g. 2827.83"),
    ("odometer_reported", False,
     "Odometer at the pump. Leave BLANK if not captured — do not guess. "
     "A reading implying more than ~3,000 km since this vehicle's last "
     "fill will be flagged for review rather than accepted automatically "
     "— check the date and the reading are both correct for large gaps "
     "between fills."),
    ("reference_number", False,
     "Receipt / OR number. Used to avoid double-counting a re-imported "
     "statement."),
    ("driver_name", False, "For reference only."),
]


import re as _re


def _norm(value):
    """Lowercase, strip, and collapse punctuation to spaces.

    The exact-match approach this replaces failed on real, ordinary
    headers -- including this app's OWN export column names, like
    "Odometer (reported)" and "Price/L" -- because neither the
    parentheses nor the slash were stripped before comparing against a
    plain-text alias. A natural round-trip (export a file, edit it,
    re-import it) or a real supplier statement with slightly different
    punctuation would otherwise silently map nothing and every reading
    would come in as MISSING, which is exactly what happened importing
    a real sample file before this fix.
    """
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = _re.sub(r"[^a-z0-9]+", " ", text)
    return _re.sub(r"\s+", " ", text).strip()


def _norm_plate(value):
    """Plate/conduction comparison, stricter than _norm(): every
    non-alphanumeric character is removed rather than collapsed to a
    space, so "AHA-9195", "AHA 9195" and "AHA9195" all compare equal.

    A provider statement (Petron's included) often has no separator at
    all in its own plate column while FMS stores plates WITH one, so
    _norm() alone -- which turns a dash into a space -- is not enough:
    "aha 9195" (from "AHA-9195") would still not equal "aha9195" (from
    "AHA9195")."""
    return _re.sub(r"[^a-z0-9]", "", _norm(value))


def _looks_like_html(raw: bytes) -> bool:
    head = raw[:1024].strip().lower()
    return head.startswith(b"<html") or head.startswith(b"<!doctype") \
        or b"<table" in head[:200]


def _parse_html_table_rows(raw: bytes):
    """Some fuel-card providers (Petron's own portal export included)
    hand out a statement whose file EXTENSION says .xls but whose
    actual content is an HTML <table> -- not a binary workbook at all.
    openpyxl rejects that outright, so the real content is detected
    first and read with a plain HTML parser instead."""
    soup = BeautifulSoup(raw, "html.parser")
    table = soup.find("table")
    if table is None:
        raise ValueError(
            "This looks like an HTML file, but no table was found in it.")
    rows = []
    for tr in table.find_all("tr"):
        cells = [cell.get_text(strip=True)
                for cell in tr.find_all(["td", "th"])]
        if any(c != "" for c in cells):
            rows.append(cells)
    if not rows:
        raise ValueError("The file's table has no rows.")
    return rows


def _load_statement_rows(file_stream):
    """Return (header_row, data_rows), detecting an HTML-table
    statement by its actual content rather than trusting the file's
    extension. Everything else still goes through openpyxl exactly as
    before."""
    raw = file_stream.read()
    if not raw or not raw.strip():
        raise ValueError("The file is empty.")
    if _looks_like_html(raw):
        rows = _parse_html_table_rows(raw)
        return rows[0], rows[1:]
    try:
        wb = load_workbook(BytesIO(raw), data_only=True)
    except Exception as exc:
        raise ValueError(
            f"Could not read this file as an Excel workbook: {exc}")
    ws = wb["Fuel"] if "Fuel" in wb.sheetnames else wb[wb.sheetnames[0]]
    all_rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    if not all_rows:
        raise ValueError("The file has no rows.")
    return all_rows[0], all_rows[1:]


def map_headers(header_row):
    """Map a statement's own headers onto our field names.

    Returns {field_name: column_index}. Unrecognised columns are simply
    ignored rather than failing the import — a provider statement
    carries plenty of columns we have no use for, and refusing the file
    over them would be obstructive.

    Both the header text and the alias list are run through the same
    _norm(), so punctuation on either side (a header's parentheses, an
    alias's own period or slash) can't cause a false mismatch.
    """
    normalized_aliases = {
        field: {_norm(a) for a in aliases}
        for field, aliases in COLUMN_ALIASES.items()
    }
    mapping = {}
    for index, raw in enumerate(header_row):
        text = _norm(raw)
        if not text:
            continue
        for field, aliases in normalized_aliases.items():
            if field in mapping:
                continue
            if text in aliases:
                mapping[field] = index
                break
    return mapping


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Fuel"
    ws.append([c for c, _, _ in TEMPLATE_COLUMNS])
    fill = PatternFill("solid", fgColor="1F3B4D")
    for i, (name, _r, _n) in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        ws.column_dimensions[cell.column_letter].width = max(16, len(name) + 4)

    ws.append(["2026-01-15 09:30", "ABC-1234", "7001234567890", "Petron EDSA",
              "DIESEL", 45.5, 62.15, 2827.83, 65430, "OR-100234", "Juan Cruz"])
    for i in range(1, len(TEMPLATE_COLUMNS) + 1):
        ws.cell(row=2, column=i).font = Font(italic=True, color="888888")
    ws.cell(row=3, column=1).value = "^ Delete this example row before uploading."
    ws.cell(row=3, column=1).font = Font(italic=True, bold=True, color="C00000")

    ref = wb.create_sheet("Reference")
    ref.append(["Column", "Required", "Notes"])
    for c in ref[1]:
        c.font = Font(bold=True)
    for name, required, note in TEMPLATE_COLUMNS:
        ref.append([name, "YES" if required else "optional", note])
    ref.column_dimensions["A"].width = 24
    ref.column_dimensions["C"].width = 70
    ref.append([])
    ref.append(["A provider statement can be uploaded directly — columns are "
               "matched by name, so exact ordering does not matter."])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _to_decimal(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value).replace(",", "").replace("₱", "").strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f"'{value}' is not a valid number")


def _to_int(value):
    d = _to_decimal(value)
    return int(d) if d is not None else None


def _to_datetime(value):
    if value is None or str(value).strip() == "":
        raise ValueError("a transaction date is required")
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M",   # <input type="datetime-local">
               "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
               "%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %I:%M %p",  # Petron: 9/30/2026 1:46:08 PM
               "%m/%d/%Y %H:%M", "%m/%d/%Y", "%d/%m/%Y %H:%M", "%d/%m/%Y",
               "%d-%b-%Y %H:%M", "%d-%b-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError(f"'{text}' is not a recognisable date")


def import_fuel(file_stream, dry_run: bool = True, filename: str = None,
                user_id: int = None) -> dict:
    """Load a fuel statement -- per-row validation, then a single
    commit decision for the whole file.

    The commit MODE is a SystemParameter (`FUEL_IMPORT_MODE`), not a
    hardcoded rule:
      * ALL_OR_NOTHING (default) -- if any row has a genuine validation
        error (bad date, missing required value, no matching vehicle),
        NOTHING is saved. A duplicate is not treated as that kind of
        error: it is an expected, defensive skip, and does not by
        itself block the rest of a re-uploaded statement.
      * PARTIAL -- the previous behaviour: a bad row is skipped and
        reported, the rest still import. Kept available because an
        all-or-nothing failure over one bad cell in a 200-row monthly
        statement is not free either; some clients may prefer it.

    Either way, every row's problems are collected up front (same as a
    dry-run preview) so the result always carries a clear, specific
    summary of what went wrong -- not just "nothing happened".

    Every row actually created is stamped with the same
    import_batch_id, the original filename, who imported it, and when
    -- so a wrong upload can be found and undone as a whole rather than
    row by row.
    """
    import uuid
    from datetime import datetime as _dt
    from app.modules.master_data.vehicle.models import Vehicle
    from app.modules.transactions.fuel.models import (
        FuelCard, FuelTransaction)
    from app.modules.transactions.fuel.odometer_validation import (
        OdometerValidationService)
    from app.modules.transactions.fuel.analytics import FuelAnalyticsService
    from app.modules.system_admin.services.system_parameter_service import (
        SystemParameterService)

    batch_id = uuid.uuid4().hex[:12]
    imported_at = _dt.now()
    mode = SystemParameterService().get("FUEL_IMPORT_MODE", "ALL_OR_NOTHING")

    header, data_rows = _load_statement_rows(file_stream)
    mapping = map_headers(header)

    missing = [f for f in ("transaction_date", "litres", "total_amount")
              if f not in mapping]
    if missing:
        raise ValueError(
            "Could not find these column(s) in the file: "
            + ", ".join(missing)
            + ". Columns are matched by name — check the header row, or "
              "start from the downloaded template.")
    if "plate_number" not in mapping and "card_number" not in mapping:
        raise ValueError(
            "The file needs either a plate/vehicle column or a card number "
            "column, so each fill can be matched to a vehicle.")

    vehicles = {}
    for v in Vehicle.query.all():
        for key in (v.plate_number, v.conduction_number):
            if key:
                vehicles[_norm_plate(key)] = v
    cards = {_norm(c.card_number): c for c in FuelCard.query.all()
            if c.card_number}

    stats = {"total_rows": 0, "created": 0, "skipped": 0, "errors": []}
    validator = OdometerValidationService()
    analytics = FuelAnalyticsService()
    ready_rows = []          # rows that passed validation, not yet saved
    seen_in_file = set()     # (vehicle_id, when, reference) within THIS file

    for row_number, row in enumerate(data_rows, start=2):
        if row is None or all(_norm(v) == "" for v in row):
            continue

        def get(field):
            index = mapping.get(field)
            if index is None or index >= len(row):
                return None
            return row[index]

        plate = _norm_plate(get("plate_number"))
        if plate == "abc1234":
            continue                     # the template's example row

        stats["total_rows"] += 1
        errors = []

        vehicle = vehicles.get(plate) if plate else None
        card = cards.get(_norm(get("card_number")))
        if vehicle is None and card is not None:
            # A card assigned to a vehicle identifies the vehicle even
            # when the statement doesn't name one.
            vehicle = card.vehicle
        if vehicle is None:
            errors.append(
                f"no vehicle matches '{get('plate_number') or get('card_number')}'")

        try:
            when = _to_datetime(get("transaction_date"))
        except ValueError as exc:
            errors.append(f"transaction_date: {exc}")
            when = None

        try:
            litres = _to_decimal(get("litres"))
            if litres is None or litres <= 0:
                errors.append("litres must be greater than zero")
        except ValueError as exc:
            errors.append(f"litres: {exc}")
            litres = None

        try:
            amount = _to_decimal(get("total_amount"))
            if amount is None:
                errors.append("total_amount is required")
        except ValueError as exc:
            errors.append(f"total_amount: {exc}")
            amount = None

        try:
            price = _to_decimal(get("price_per_litre"))
        except ValueError:
            price = None                 # derived below rather than failing
        try:
            odometer = _to_int(get("odometer_reported"))
        except ValueError:
            # A junk odometer must NOT fail the row -- the fill still
            # happened and the money is real. It is imported without a
            # reading and shows as MISSING, which is exactly what the
            # validation layer is for.
            odometer = None

        reference = get("reference_number")
        reference = str(reference).strip() if reference else None

        if errors:
            stats["skipped"] += 1
            stats["errors"].append({
                "row": row_number,
                "identifier": get("plate_number") or get("card_number") or "(blank)",
                "problems": errors, "type": "error"})
            continue

        if price is None and litres and amount:
            price = round(amount / litres, 4)

        dup_key = (vehicle.id, when, reference)
        already_in_db = FuelTransaction.query.filter_by(
            vehicle_id=vehicle.id, transaction_date=when,
            reference_number=reference).first() is not None
        if already_in_db or dup_key in seen_in_file:
            stats["skipped"] += 1
            stats["errors"].append({
                "row": row_number, "identifier": vehicle.plate_number,
                "problems": ["already imported (same vehicle, date and "
                            "reference) — skipped to avoid double-counting"],
                "type": "duplicate"})
            continue
        seen_in_file.add(dup_key)

        ready_rows.append({
            "row_number": row_number, "vehicle": vehicle, "card": card,
            "when": when, "station": get("station"),
            "fuel_type": get("fuel_type"), "litres": litres, "price": price,
            "amount": amount, "odometer": odometer, "reference": reference,
        })
        stats["created"] += 1

    blocking_errors = [e for e in stats["errors"] if e.get("type") != "duplicate"]

    if dry_run:
        stats["committed"] = False
        stats["import_batch_id"] = None
    elif mode == "ALL_OR_NOTHING" and blocking_errors:
        # Nothing in this file is saved: a partial save here would
        # leave the person unable to tell, from the result alone,
        # which rows actually went in without re-checking the whole
        # file against the database.
        stats["created"] = 0
        stats["committed"] = False
        stats["import_batch_id"] = None
    else:
        created_rows = []
        for r in ready_rows:
            txn = FuelTransaction(
                vehicle_id=r["vehicle"].id,
                fuel_card_id=r["card"].id if r["card"] else None,
                transaction_date=r["when"],
                station=str(r["station"]).strip() if r["station"] else None,
                fuel_type=(str(r["fuel_type"]).strip().upper()
                          if r["fuel_type"] else None),
                litres=r["litres"], price_per_litre=r["price"],
                total_amount=r["amount"], odometer_reported=r["odometer"],
                reference_number=r["reference"], source="IMPORT",
                import_filename=filename, import_batch_id=batch_id,
                imported_by=user_id, imported_at=imported_at)
            db.session.add(txn)
            db.session.flush()
            created_rows.append(txn)

        if created_rows:
            # Validate oldest-first per vehicle, so each fill measures
            # from a predecessor that has already been assessed.
            # Validating in file order would compare against readings
            # not yet checked.
            for vehicle_id in {t.vehicle_id for t in created_rows}:
                validator.revalidate_vehicle(vehicle_id)
            for txn in created_rows:
                analytics.detect_anomalies(txn, commit=False)
            db.session.commit()

        stats["flagged"] = sum(1 for t in created_rows if t.anomaly_flags)
        stats["untrusted_odometer"] = sum(
            1 for t in created_rows
            if t.odometer_status in ("SUSPECT", "MISSING"))
        stats["committed"] = True
        stats["import_batch_id"] = batch_id if created_rows else None

    stats.setdefault("flagged", 0)
    stats.setdefault("untrusted_odometer", 0)
    stats["filename"] = filename

    if dry_run or not stats["committed"]:
        if blocking_errors:
            stats["message"] = (
                f"{len(blocking_errors)} of {stats['total_rows']} row(s) "
                "have errors. " +
                ("Nothing was imported — fix the issues below and "
                 "re-upload." if not dry_run else
                 "Fix the issues below, then import.")
            )
        elif not dry_run:
            stats["message"] = (
                "Nothing in this file is new — every row was already "
                "imported previously." if stats["skipped"] and not
                stats["created"] else None)
    return stats


def list_import_batches(limit=30):
    """Recent import batches, most recent first, for the batch filter
    dropdown and the "undo a bad upload" screen."""
    from app.modules.transactions.fuel.models import FuelTransaction
    from app.modules.user_management.models import User

    rows = (db.session.query(
               FuelTransaction.import_batch_id,
               FuelTransaction.import_filename,
               FuelTransaction.imported_by,
               FuelTransaction.imported_at,
               func.count(FuelTransaction.id).label("row_count"),
               func.sum(FuelTransaction.total_amount).label("total_spend"))
           .filter(FuelTransaction.import_batch_id.isnot(None))
           .group_by(FuelTransaction.import_batch_id,
                    FuelTransaction.import_filename,
                    FuelTransaction.imported_by,
                    FuelTransaction.imported_at)
           .order_by(FuelTransaction.imported_at.desc())
           .limit(limit).all())

    users = {u.id: u for u in User.query.all()}
    return [{
        "batch_id": r.import_batch_id, "filename": r.import_filename,
        "imported_by": users.get(r.imported_by),
        "imported_at": r.imported_at, "row_count": r.row_count,
        "total_spend": r.total_spend or 0,
    } for r in rows]


def delete_import_batch(batch_id):
    """Remove every transaction from one upload -- the undo for a wrong
    file. Deliberately keyed on the batch, not a date range or a vehicle
    list, so it removes exactly and only what one upload added, nothing
    from before it and nothing added since by a different upload."""
    from app.modules.transactions.fuel.models import FuelTransaction
    rows = FuelTransaction.query.filter_by(import_batch_id=batch_id).all()
    count = len(rows)
    affected_vehicles = {r.vehicle_id for r in rows}
    for row in rows:
        db.session.delete(row)
    db.session.commit()

    # Deleting a batch can remove the reading a later, still-present fill
    # was measuring FROM -- that later fill's distance/km-L would then be
    # silently wrong until recomputed against whatever baseline remains.
    from app.modules.transactions.fuel.odometer_validation import (
        OdometerValidationService)
    validator = OdometerValidationService()
    for vehicle_id in affected_vehicles:
        validator.revalidate_vehicle(vehicle_id)

    return count


def export_fuel(rows) -> bytes:
    """Export transactions, including the validation verdict.

    The odometer status and note travel with the data deliberately: a
    spreadsheet showing consumption without showing which rows were
    excluded from it would invite exactly the false confidence this
    module exists to prevent.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Fuel Transactions"
    headers = ["Date", "Vehicle", "Card", "Station", "Fuel Type", "Litres",
              "Price/L", "Total", "Odometer (reported)", "Odometer (used)",
              "Odometer Status", "Distance (km)", "km/L", "Cost/km",
              "Flags", "Note"]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1F3B4D")
    for i in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["P"].width = 60

    for t in rows:
        ws.append([
            t.transaction_date.strftime("%Y-%m-%d %H:%M") if t.transaction_date else "",
            (t.vehicle.plate_number or t.vehicle.conduction_number) if t.vehicle else "",
            t.fuel_card.card_number if t.fuel_card else "",
            t.station or "", t.fuel_type or "",
            float(t.litres) if t.litres is not None else "",
            float(t.price_per_litre) if t.price_per_litre is not None else "",
            float(t.total_amount) if t.total_amount is not None else "",
            t.odometer_reported if t.odometer_reported is not None else "",
            t.odometer_used if t.odometer_used is not None else "",
            t.odometer_status or "",
            t.distance_km if t.distance_km is not None else "",
            float(t.km_per_litre) if t.km_per_litre is not None else "",
            float(t.cost_per_km) if t.cost_per_km is not None else "",
            t.anomaly_flags or "", t.odometer_note or "",
        ])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
