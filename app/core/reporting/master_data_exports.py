"""Master data exports.

Purpose: let someone (often the client, during configuration review)
pull out exactly what is set up in the system, without reading it off
a paginated screen a page at a time.

Built as ONE generic generator driven by a per-module column
definition, rather than a dozen near-identical functions. Adding a
module means adding a few lines to MASTER_DATA_EXPORTS below, not
writing another export.

The PM Scope export is the reason this exists and is deliberately the
richest: it exports every scope template AND every activity line
underneath it, on two sheets, because "what is configured" for PM
scope means the activities, not just the template names.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.extensions import db

_HEADER_FILL = PatternFill("solid", fgColor="1F3B4D")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb, name, title, headers, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name[:31]          # Excel's own limit
    ws.append([title])
    ws["A1"].font = Font(size=14, bold=True)
    ws.append([])
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(3, c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = max(
            14, min(len(str(headers[c - 1])) + 4, 50))
    ws.freeze_panes = "A4"
    return ws


def _value(obj, path):
    """Read a possibly-nested attribute ("branch.name") safely.

    Returns an em dash rather than raising when any link in the chain is
    missing -- a vehicle type with no category shouldn't break the whole
    export of every other row.
    """
    cur = obj
    for part in path.split("."):
        if cur is None:
            return "—"
        cur = getattr(cur, part, None)
    if cur is None or cur == "":
        return "—"
    if cur is True:
        return "Yes"
    if cur is False:
        return "No"
    return cur


def generate_master_data_xlsx(key: str) -> tuple:
    """(bytes, filename) for one master data module.

    Raises KeyError for an unknown key so a bad URL is a clean 404
    rather than an empty spreadsheet that looks like "nothing is
    configured".
    """
    spec = MASTER_DATA_EXPORTS[key]
    wb = Workbook()

    rows = spec["query"]()
    ws = _sheet(wb, spec["sheet"], spec["title"], spec["headers"], first=True)
    for obj in rows:
        ws.append([_value(obj, f) if isinstance(f, str) else f(obj)
                  for f in spec["fields"]])

    # Optional second sheet for the detail lines under each record.
    if spec.get("detail"):
        d = spec["detail"]
        ws2 = _sheet(wb, d["sheet"], d["title"], d["headers"])
        for obj in rows:
            for line in d["lines"](obj):
                ws2.append([f(obj, line) for f in d["fields"]])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), spec["filename"]


# ── Per-module definitions ──────────────────────────────────────────────
#
# Each entry is what a person reviewing the configuration would need to
# see. Deliberately NOT a dump of every database column: internal ids,
# audit timestamps and soft-delete flags are noise in a configuration
# review and would obscure the fields that actually matter.

def _pm_scope_query():
    from app.modules.maintenance_config.models import PMScopeTemplate
    from sqlalchemy.orm import joinedload, selectinload
    return (PMScopeTemplate.query
           .options(joinedload(PMScopeTemplate.maintenance_type),
                    joinedload(PMScopeTemplate.pm_schedule),
                    selectinload(PMScopeTemplate.items))
           .order_by(PMScopeTemplate.name).all())


MASTER_DATA_EXPORTS = {
    "pm-scope-templates": {
        "title": "PM Scope Templates — System Configuration",
        "filename": "PM_Scope_Templates.xlsx",
        "sheet": "Scope Templates",
        "query": _pm_scope_query,
        "headers": ["Template Name", "Maintenance Type",
                    "Linked PM Template", "Interval (km)",
                    "Interval (days)", "Activities", "Active"],
        "fields": [
            "name",
            "maintenance_type.name",
            lambda t: (t.pm_schedule.profile_description
                      if t.pm_schedule else "—"),
            lambda t: (t.pm_schedule.interval_km
                      if t.pm_schedule and t.pm_schedule.interval_km else "—"),
            lambda t: (t.pm_schedule.interval_days
                      if t.pm_schedule and t.pm_schedule.interval_days else "—"),
            lambda t: len(t.items),
            lambda t: "Yes" if t.is_active else "No",
        ],
        # The activities ARE the configuration for PM scope -- a list of
        # template names alone would not answer "what does this PM
        # actually cover", which is the question being asked.
        "detail": {
            "sheet": "Scope Activities",
            "title": "PM Scope Activities — every line under every template",
            "headers": ["Template Name", "Maintenance Type", "#",
                        "Activity Code", "Activity Description",
                        "Labor Hours", "Est. Cost", "Required Parts"],
            "lines": lambda t: sorted(t.items, key=lambda i: i.sort_order),
            "fields": [
                lambda t, i: t.name,
                lambda t, i: (t.maintenance_type.name
                             if t.maintenance_type else "—"),
                lambda t, i: i.sort_order + 1,
                lambda t, i: i.activity_code or "—",
                lambda t, i: i.activity_description,
                lambda t, i: i.standard_labor_hours or "—",
                lambda t, i: i.estimated_cost or "—",
                lambda t, i: i.required_parts or "—",
            ],
        },
    },
}


def _register(key, *, title, filename, sheet, model_path, headers, fields,
             order_by=None):
    """Register a straightforward single-sheet master data export."""
    def _q():
        import importlib
        module_name, cls_name = model_path.rsplit(".", 1)
        model = getattr(importlib.import_module(module_name), cls_name)
        q = model.query
        if order_by:
            q = q.order_by(getattr(model, order_by))
        return q.all()
    MASTER_DATA_EXPORTS[key] = {
        "title": title, "filename": filename, "sheet": sheet,
        "query": _q, "headers": headers, "fields": fields,
    }


_register(
    "pm-schedules",
    title="PM Templates (Schedules) — System Configuration",
    filename="PM_Templates.xlsx", sheet="PM Templates",
    model_path="app.modules.maintenance_config.models.PMSchedule",
    order_by="profile_code",
    headers=["Profile Code", "Description", "Maintenance Type", "Package #",
             "Trigger", "Interval (km)", "Interval (days)",
             "Interval (hours)", "Cumulative km", "Make", "Model",
             "Vehicle Type", "Active"],
    fields=["profile_code", "profile_description", "maintenance_type.name",
            "sequence_position", "trigger_mode", "interval_km",
            "interval_days", "interval_hours", "cumulative_km",
            lambda s: (s.vehicle_brand.name if s.vehicle_brand
                      else (s.vehicle_make or "—")),
            lambda s: (s.vehicle_model_ref.name if s.vehicle_model_ref
                      else (s.vehicle_model or "—")),
            "vehicle_type.name",
            lambda s: "Yes" if s.is_active else "No"])

_register(
    "vehicles",
    title="Vehicle Master — System Configuration",
    filename="Vehicle_Master.xlsx", sheet="Vehicles",
    model_path="app.modules.master_data.vehicle.models.Vehicle",
    order_by="plate_number",
    headers=["Plate No.", "Conduction No.", "Brand", "Model", "Year",
             "Vehicle Type", "Branch", "Department", "Cost Center",
             "Status", "Odometer", "Active"],
    fields=["plate_number", "conduction_number", "brand", "model", "year",
            "vehicle_type.name", "branch.name", "department.name",
            "cost_center", "status", "current_odometer",
            lambda v: "Yes" if v.is_active else "No"])

_register(
    "branches",
    title="Branches — System Configuration",
    filename="Branches.xlsx", sheet="Branches",
    model_path="app.modules.master_data.org.models.Branch",
    order_by="code",
    headers=["Code", "Name", "Address", "Active"],
    fields=["code", "name", "address",
            lambda b: "Yes" if b.is_active else "No"])

_register(
    "departments",
    title="Departments — System Configuration",
    filename="Departments.xlsx", sheet="Departments",
    model_path="app.modules.master_data.org.models.Department",
    order_by="code",
    headers=["Code", "Name", "Branch", "Cost Center", "Description",
             "Active"],
    fields=["code", "name", "branch.name", "cost_center", "description",
            lambda d: "Yes" if d.is_active else "No"])

_register(
    "vehicle-types",
    title="Vehicle Types — System Configuration",
    filename="Vehicle_Types.xlsx", sheet="Vehicle Types",
    model_path="app.modules.master_data.reference.models.VehicleType",
    order_by="code",
    headers=["Code", "Name", "Category", "Active"],
    fields=["code", "name", "category",
            lambda t: "Yes" if t.is_active else "No"])

_register(
    "maintenance-types",
    title="Maintenance Types — System Configuration",
    filename="Maintenance_Types.xlsx", sheet="Maintenance Types",
    model_path="app.modules.master_data.reference.models.MaintenanceType",
    order_by="code",
    headers=["Code", "Name", "Category", "Active"],
    fields=["code", "name", "category",
            lambda t: "Yes" if t.is_active else "No"])

_register(
    "drivers",
    title="Drivers / Assignees — System Configuration",
    filename="Drivers_Assignees.xlsx", sheet="Drivers",
    model_path="app.modules.master_data.driver.models.Driver",
    order_by="employee_number",
    headers=["Employee No.", "Person ID", "First Name", "Last Name",
             "Assignee Type", "Job Title", "Branch", "License No.",
             "License Type", "License Expiry", "Has Photo", "Active"],
    fields=["employee_number", "person_id", "first_name", "last_name",
            "assignee_type", "job_title", "branch.name", "license_number",
            "license_type", "license_expiry",
            lambda d: "Yes" if d.photo_attachment_id else "No",
            lambda d: "Yes" if d.is_active else "No"])

_register(
    "vendors",
    title="Vendors — System Configuration",
    filename="Vendors.xlsx", sheet="Vendors",
    model_path="app.modules.master_data.vendor.models.Vendor",
    order_by="code",
    headers=["Code", "Name", "Contact Person", "Phone", "Email", "Address",
             "Active"],
    fields=["code", "name", "contact_person", "phone", "email", "address",
            lambda v: "Yes" if v.is_active else "No"])

_register(
    "vehicle-brands",
    title="Vehicle Brands — System Configuration",
    filename="Vehicle_Brands.xlsx", sheet="Brands",
    model_path="app.modules.master_data.vehicle_brand.models.VehicleBrand",
    order_by="name",
    headers=["Name", "Models", "Active"],
    fields=["name", lambda b: len(b.models),
            lambda b: "Yes" if b.is_active else "No"])

_register(
    "vehicle-models",
    title="Vehicle Models — System Configuration",
    filename="Vehicle_Models.xlsx", sheet="Models",
    model_path="app.modules.master_data.vehicle_brand.models.VehicleModel",
    order_by="name",
    headers=["Brand", "Model", "Active"],
    fields=["brand.name", "name",
            lambda m: "Yes" if m.is_active else "No"])
